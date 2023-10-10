import tkinter as tk
from tkinter import filedialog, ttk, Listbox, Scrollbar
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from datetime import datetime
from pptx import Presentation
from pptx.util import Cm
from pathlib import Path
import re
import os

data_frames = {}
red_dot = None
data_groups = {}

root = tk.Tk()
root.title("Load Regulation_1.0.1") # windows title / PGM version

def analyze_filename(filename):
    match = re.search(r'(\w)(\d)(\d)[pP](\d)(\d+)', filename) # analyze file name(pin name / voltage / temperature) 

    if match:
        pin_name = match.group(1) # LDO or BUCK
        pin_no = match.group(2) # pin no
        voltage = float(match.group(3) + '.' + match.group(4)) # convert p to .
        temperature = int(match.group(5)) # temperature
        
        if temperature == 20: # convert 20 to -20
            temperature = "-20"
    else:
        pin_name = None
        pin_no + None
        voltage = None
        temperature = None
    return pin_name, pin_no, voltage, temperature

def open_file():
    global data_frames, red_dot, data_groups

    file_paths = filedialog.askopenfilenames(title="Select CSV files", filetypes=[("CSV files", "*.csv")]) # data file path

    if not file_paths:
        print("Please select CSV files")
    else:
        ax.clear()
        data_frames = {}
        data_groups = {}
        
        if red_dot is not None: # clear red dot
            red_dot.remove()
            red_dot = None

        for file_path in file_paths:
            file_name = os.path.splitext(os.path.basename(file_path))[0]
            df = pd.read_csv(file_path, skiprows=7) # data file in upper 7 == NULL
            df.rename(columns={'Reading': 'V', 'Value': 'mA'}, inplace=True) # analyze data file
            df['mA'] = df['mA'] * -1000 # calculate
            data_frames[file_name] = df
            pin_name, pin_no, voltage, temperature = analyze_filename(file_name)
            
            # convert pin name clearly
            if pin_name == "l":
                pin_name = "LDO"
            elif pin_name == "b":
                pin_name = "BUCK"
            else:
                pin_name = None

            if pin_name is not None and pin_no is not None and temperature is not None:
                data_group_name = f"{pin_name}{pin_no}_{temperature}°C" # data group name
                if data_group_name not in data_groups:
                    data_groups[data_group_name] = []
                data_groups[data_group_name].append(file_name)

            ax.plot(df['mA'], df['V'], linestyle='-', label=file_name) # show all graph

        pin_name, pin_no, voltage, temperature = analyze_filename(list(data_frames.keys())[0]) # graph option
        title = "Total Data Plot"

        ax.set_title(title)
        ax.set_xlabel('Current (mA)')
        ax.set_ylabel('Voltage (V)')
        ax.grid(True)

        fig.canvas.mpl_connect('button_press_event', on_click) # mouse event function
        info_box.set("")
        update_listbox()
        canvas.draw()

def update_listbox():
    listbox.delete(0, tk.END)

    for data_group_name in data_groups:
        listbox.insert(tk.END, data_group_name)

def on_click(event): # mouse event function
    global red_dot 

    if event.button == 1: # when left button clicked
        x_value, y_value = event.xdata, event.ydata
        closest_point_info = None
        closest_distance = float('inf')
        
        selected_index = listbox.curselection()
        if selected_index:
            selected_index = int(selected_index[0])
            selected_data_group = listbox.get(selected_index)
            
            for file_name in data_groups[selected_data_group]: # to find real data when mouse clicked
                df = data_frames[file_name]
                if df['mA'].min() <= x_value <= df['mA'].max() and df['V'].min() <= y_value <= df['V'].max():
                    distances = ((df['mA'] - x_value) ** 2 + (df['V'] - y_value) ** 2)
                    min_distance_idx = distances.idxmin()
                    if distances[min_distance_idx] < closest_distance:
                        closest_point_info = {
                            'file_name': file_name,
                            'idx': min_distance_idx
                        }
                        closest_distance = distances[min_distance_idx]

        if closest_point_info: # clicked point info
            file_name = closest_point_info['file_name']
            closest_point_idx = closest_point_info['idx']

            if red_dot is not None:
                red_dot.remove()

            # clicked point info
            red_dot = ax.plot(data_frames[file_name]['mA'][closest_point_idx], data_frames[file_name]['V'][closest_point_idx], 'ro', markersize=6)[0]
            canvas.draw()
            info_box.set(f"Clicked Point\n[File Name: {file_name}]\n\nCurrent: {data_frames[file_name]['mA'][closest_point_idx]:.2f} mA\nVoltage: {data_frames[file_name]['V'][closest_point_idx]:.2f} V")

def on_listbox_select(event):
    selected_index = listbox.curselection()
    if selected_index:
        selected_index = int(selected_index[0])
        selected_data_group = listbox.get(selected_index)

        for file_name in data_groups[selected_data_group]: # to check data file name
            print("selected file: " + file_name)
        print("*************************")

        ax.clear()
        for file_name in data_groups[selected_data_group]:
            df = data_frames[file_name]
            ax.plot(df['mA'], df['V'], linestyle='-', label = file_name[2] + "." + file_name[4] + "V") # legends = voltage

        ax.set_title(selected_data_group)
        ax.set_xlabel('Current (mA)')
        ax.set_ylabel('Voltage (V)')
        ax.legend(loc='upper center', bbox_to_anchor=(0.1, 1.0), fancybox=True, ncol=3)
        ax.grid(True)
        canvas.draw()

def save_as_png():
    save_folder = filedialog.askdirectory(title="Select a folder to save the PNG file")
    if save_folder:
        for data_group_name in data_groups:
            fig, ax = plt.subplots(figsize=(10, 6))
            for file_name in data_groups[data_group_name]:
                df = data_frames[file_name]
                ax.plot(df['mA'], df['V'], linestyle='-', label = file_name[2] + "." + file_name[4] + "V") # legends = voltage
                
            ax.set_title(data_group_name) # graph option
            ax.grid(True)
            ax.set_xlabel('Current (mA)')
            ax.set_ylabel('Voltage (V)')
            ax.legend(loc='upper center', bbox_to_anchor=(0.5, -0.1), fancybox=True, ncol=3)

            fixed_file_name = f"{data_group_name}.png" # file name, path
            png_file_path = os.path.join(save_folder, fixed_file_name)

            fig.savefig(png_file_path, bbox_inches='tight') # save as png
            print(f"Save PNG path: {png_file_path}")

def create_ppt():
    png_files = filedialog.askopenfilenames(title="Select PNG files", filetypes=[("PNG files", "*.png")])

    if not png_files:
        print("No PNG files selected.")
        return

    grouped_files = {}

    for png_file in png_files:
        file_name = os.path.basename(png_file)
        parts = file_name.split('_')
        group_name = '_'.join(parts[:1]) # separate dat name by _
        grouped_files.setdefault(group_name, []).append(png_file)

    template_path = "JCET_Format.pptx"  # ppt template
    presentation = Presentation(template_path)

    for group_name, files in grouped_files.items():
        slide = presentation.slides.add_slide(presentation.slide_layouts[1]) # ppt templat page 2(1 == title)

        title = slide.shapes.title
        title.text = group_name

        img_width = Cm(10.55) # png image size in ppt
        img_height = Cm(7.04)
        start_x = Cm(0.96) # start point

        for i, file in enumerate(files):
            left = start_x + i * img_width

            if i >= 3: # second row in ppt page
                top = Cm(10.06)
                start_x = Cm(0.96)
            else: # first row in ppt page
                top = Cm(3.36)
            
            slide.shapes.add_picture(file, left, top, width=img_width, height=img_height)

    formatted_date = datetime.today().strftime('%y%m%d_%H%M%S') # file name
    file_name = f"{formatted_date}_Load Regulation.pptx"

    save_path = filedialog.asksaveasfilename(
        title="Save PPT", 
        filetypes=[("PowerPoint files", "*.pptx")], 
        initialfile=file_name
    )

    if save_path:
        presentation.save(save_path)
        print(f"Presentation saved path: {save_path}")

# main
fig, ax = plt.subplots(figsize=(10, 6)) # graph size

open_button = ttk.Button(root, text="Open", command=open_file) # open
open_button.pack(side=tk.TOP, padx=10, pady=5)

save_button = ttk.Button(root, text="Save png", command=save_as_png) # save
save_button.pack(side=tk.TOP, padx=10, pady=5)

ppt_button = ttk.Button(root, text="Create PPT", command=create_ppt)
ppt_button.pack(side=tk.TOP, padx=10, pady=5)

quit_button = ttk.Button(root, text="Quit", command=root.quit) # quit
quit_button.pack(side=tk.TOP, padx=10, pady=5)

canvas = FigureCanvasTkAgg(fig, master=root)
canvas_widget = canvas.get_tk_widget()
canvas_widget.pack(side=tk.LEFT, fill=tk.BOTH, expand=1)

info_box = tk.StringVar()
info_label = tk.Label(root, textvariable=info_box, padx=10, pady=5, borderwidth=10) # clicked point
info_label.pack(side=tk.TOP)

listbox_frame = tk.Frame(root)
listbox_frame.pack(side=tk.LEFT, padx=10, pady=5)

listbox_label = tk.Label(listbox_frame, text="Data Groups") # data list
listbox_label.pack()

listbox = Listbox(listbox_frame, selectmode=tk.SINGLE)
listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=1)

scrollbar = Scrollbar(listbox_frame, orient=tk.VERTICAL) # data list scrollbar
scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

listbox.config(yscrollcommand=scrollbar.set)
scrollbar.config(command=listbox.yview)

update_listbox()
listbox.bind("<<ListboxSelect>>", on_listbox_select)

root.mainloop()

# 231006
# legends changed at saved png