"""Consolidate everything the reviewer (LEE SeokJin) asked for into the ONE
existing workbook (report/Harald_260820/<date>_Harald_260820_Summary.xlsx),
per user confirmation 2026-08-20:

1. Add one "<Group> Summary" sheet per top-level Bench subfolder -- same
   Mean/Std/Min/Max/USL/LSL + Delta shape as the existing pooled
   "Harald_260820 Summary" sheet (kept as-is, untouched), but computed from
   ONLY that group's DUTs (reuses build_harald_260820_xlsx_summary_by_
   vendor.py's row-building, imported not modified). Originally exactly 2
   groups (Skyworks/Sony); 2026-08-25 generalized to however many top-level
   subfolders Bench actually has (e.g. a 100-DUT pull's DNNJ/DSKJ/SNKJ/
   SSNJ, 4 groups) -- see vendor_lists_from_nested_bench's own docstring
   for why this isn't collapsed into a smaller "real vendor" count.
2. REPLACE the existing flat "Raw Values" sheet with a 3-level-header
   version grouped Bench/ATE/Delta -> group -> individual DUT serial
   ("2did") -- matching the reviewer's screenshot (Bench -> SKYWORKS ->
   DUT-serial columns) extended with ATE/Delta blocks the same way (reuses
   add_harald_260820_delta_by_vendor_sheet.py's row-building).

Group/DUT lists: originally Downloads/"260818_2did 40pcs.txt" (SKYWORKS 20
+ SONY 20), now derived straight from the Bench folder structure -- see
build_harald_260820_xlsx_summary_by_vendor.parse_vendor_lists.
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

_THIS_DIR = Path(__file__).resolve().parent
_BASE_DIR = _THIS_DIR.parent
sys.path.insert(0, str(_BASE_DIR))
sys.path.insert(0, str(_BASE_DIR / "Cal"))
sys.path.insert(0, str(_THIS_DIR))

import xlsx_summary_writer as xw
import make_cal_distribution as cal
import make_harald_260820_distribution as harald
import add_harald_260820_raw_values_sheet as raw_sheet
import build_harald_260820_xlsx_summary_by_vendor as by_vendor

import datetime
# 2026-08-21: was hardcoded to the "2026-08-20"-dated file build_harald_
# 260820_xlsx_summary.py originally wrote -- that script's own REPORT_DATE
# is today's date, so a hardcoded historical date here silently pointed at
# a stale file (or nothing, once dates drift) on any later rerun. Compute
# the same way that script does, so both always agree on today's filename.
REPORT_DATE = datetime.date.today().isoformat()
SUMMARY_PATH = harald.PATHS.report_dir / f"{REPORT_DATE}_Harald_260820_Summary.xlsx"


def write_vendor_summary_sheet(wb, sheet_name, rows, n_cols=17):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    ws.cell(row=1, column=1, value="Test Item")
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    col = 2
    for group_label, subs in (("Bench", xw.BENCH_SUB), ("ATE", xw.ATE_SUB), ("Delta", xw.DELTA_SUB)):
        start = col
        for sub in subs:
            ws.cell(row=2, column=col, value=sub)
            col += 1
        ws.cell(row=1, column=start, value=group_label)
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=col - 1)

    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=2, max_col=n_cols):
        for cell in row:
            cell.font = header_font
            cell.alignment = center
    for i in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14

    for r, row in enumerate(rows, start=3):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    ws.freeze_panes = "A3"


def main() -> None:
    print(f"Parsing vendor DUT list ({by_vendor.VENDOR_LIST_PATH.name})...")
    vendor_lists_raw = by_vendor.parse_vendor_lists(by_vendor.VENDOR_LIST_PATH)
    # 2026-08-25: was a hardcoded ("SKYWORKS", "SONY") filter -- silently
    # dropped every group when Bench's top-level subfolders are named
    # something else (a 100-DUT pull's DNNJ/DSKJ/SNKJ/SSNJ), leaving
    # vendor_order empty and crashing downstream (openpyxl merge_cells on a
    # zero-width range). Use whatever groups parse_vendor_lists actually
    # found -- sorted for a deterministic sheet/column order.
    vendor_order = sorted(vendor_lists_raw)
    vendor_lists = {v: sorted(vendor_lists_raw[v]) for v in vendor_order}
    for v in vendor_order:
        print(f"  {v}: {len(vendor_lists[v])} DUT(s)")

    print("Building hex-id -> DUT-serial map from the reference Bench copy...")
    hex_to_serial = raw_sheet.build_hex_to_serial(raw_sheet.NESTED_BENCH_DIR)
    print(f"  {len(hex_to_serial)} DUT(s) mapped")

    print("Reading Bench_Harald_260820 CSVs...")
    bench_items = cal.read_bench(harald.PATHS.bench_dir)
    print(f"  {len(bench_items)} test item(s) found")

    print(f"Parsing ATE log ({harald.ATE_PATH.name}) per-DUT...")
    ate_by_item = raw_sheet.read_ate_per_dut_auto(harald.ATE_PATH)  # ATE_PATH may be one file or a folder
    ate_items_std = cal.read_ate_auto(harald.ATE_PATH)  # ATE_PATH may be one file or a folder

    jobs, _n_no_ate = cal.build_jobs(bench_items, ate_items_std)
    print(f"  {len(jobs)} qualifying item(s)")

    print("\nOpening workbook...")
    wb = openpyxl.load_workbook(SUMMARY_PATH)

    for vendor in vendor_order:
        print(f"Building {vendor} Summary sheet...")
        rows = by_vendor.build_rows(set(vendor_lists[vendor]), hex_to_serial, bench_items, ate_by_item, jobs)
        write_vendor_summary_sheet(wb, f"{vendor} Summary", rows)

    print("Rebuilding 'Raw Values' sheet (3-level header: Bench/ATE/Delta -> vendor -> 2did)...")
    rv_rows = []
    for job in jobs:
        raw_name = job["name"]
        item = harald.clean(raw_name)
        bench_by_serial = raw_sheet.bench_values_by_serial(bench_items[raw_name], hex_to_serial)
        ate_by_serial = ate_by_item.get(raw_name, {})
        row = [item]
        for group in ("bench", "ate", "delta"):
            for vendor in vendor_order:
                for serial in vendor_lists[vendor]:
                    b = bench_by_serial.get(serial)
                    a = ate_by_serial.get(serial)
                    if group == "bench":
                        row.append(b)
                    elif group == "ate":
                        row.append(a)
                    else:
                        row.append(xw.delta(b, a))
        rv_rows.append(row)

    sheet_name = "Raw Values"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.cell(row=1, column=1, value="Test Item")
    ws.merge_cells(start_row=1, start_column=1, end_row=3, end_column=1)
    col = 2
    for group_label in ("Bench", "ATE", "Delta"):
        group_start = col
        for vendor in vendor_order:
            vendor_start = col
            for serial in vendor_lists[vendor]:
                ws.cell(row=3, column=col, value=serial)
                col += 1
            ws.cell(row=2, column=vendor_start, value=vendor)
            ws.merge_cells(start_row=2, start_column=vendor_start, end_row=2, end_column=col - 1)
        ws.cell(row=1, column=group_start, value=group_label)
        ws.merge_cells(start_row=1, start_column=group_start, end_row=1, end_column=col - 1)
    n_cols = col - 1

    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=3, max_col=n_cols):
        for cell in row:
            cell.font = header_font
            cell.alignment = center
    for i in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 12

    for r, row in enumerate(rv_rows, start=4):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    ws.freeze_panes = "B4"

    # Sheet order: pooled summary, per-group summaries (whatever groups
    # actually exist -- was hardcoded to exactly "Skyworks"/"Sony", same
    # class of bug as vendor_order above), raw values.
    desired_order = ["Harald_260820 Summary", *(f"{v} Summary" for v in vendor_order), "Raw Values"]
    wb._sheets.sort(key=lambda s: desired_order.index(s.title) if s.title in desired_order else len(desired_order))

    print(f"\nSaving {SUMMARY_PATH}...")
    wb.save(SUMMARY_PATH)
    print(f"Saved. Sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
