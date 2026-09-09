"""Add a "Raw Values" sheet to the existing Harald_260820 xlsx summary
(report/Harald_260820/2026-08-20_Harald_260820_Summary.xlsx), per user
request 2026-08-20: one row per Test Item (same set/order as that
workbook's existing summary sheet, untouched here), columns split into a
Bench block and an ATE block, each with one column per DUT -- labeled by
DUT serial number ("2did"), not an anonymous #1/#2/... index.

DUT-serial labeling ("2did"), Bench side: the flat Bench_40pcs/
Harald_260820/*.csv files this project reads are named by an internal hex
id (parsed by cal.dut_id_of from the filename, e.g. "7480083C1301C6"), not
the human-readable DUT serial ("725HW3000GR0001A6Q") ATE's own CSV header
already uses. 2026-08-20: originally recovered via a temporary reference
copy of the same retest pull at Downloads/"Bench 2/Bench", still organized
by {vendor}/{DUT serial}/*.csv. 2026-08-21: pointed at Bench_40pcs/Harald
instead -- this project's OWN original (pre-retest) vendor-nested pull,
already a real, always-present part of this workspace, so no personal-
Downloads-folder copy is needed at all; verified its hex-id set and vendor/
serial assignment are byte-for-byte identical to that Downloads copy (and
to the "260818_2did 40pcs.txt" list build_harald_260820_xlsx_summary_by_
vendor.py used to read separately -- see vendor_lists_from_nested_bench
below, which now derives that same split from here too). Either copy's
files are named IDENTICALLY to their flat-pull counterpart, so matching by
filename (via cal.dut_id_of on both) is exact regardless of which one is
in use. One file in the original Downloads copy sat three levels deep
(Sony/725HW3001L60001A6Q/725HW3001M10001A6Q/w4_pb_logfile_..._7480083C13
04DC...csv) instead of the normal two -- an apparent archive/extraction
artifact -- so this reads via Path.rglob (parent-folder name, whatever the
depth) rather than a fixed-depth glob, which picks it up correctly as its
own DUT no matter which copy is pointed at. Verified 2026-08-20: this
recovers all 40 DUTs, and the resulting serial set matches the ATE CSV's
own 40 header serials exactly (1:1, no extras on either side).

ATE side needs no external mapping -- cal.read_ate() already discards
per-DUT identity (returns a plain values list), so this file parses the
ATE CSV itself, keeping {serial: value} per Test Item, instead of importing
that function.

This only ADDS a sheet -- the existing "Bench Summary"-style first sheet
(Mean/Std/Min/Max/USL/LSL + Delta, built by build_harald_260820_xlsx_
summary.py) is left exactly as already generated, per user instruction.
"""
from __future__ import annotations

import csv
import datetime
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

_THIS_DIR = Path(__file__).resolve().parent
_BASE_DIR = _THIS_DIR.parent
sys.path.insert(0, str(_BASE_DIR / "Cal"))
sys.path.insert(0, str(_THIS_DIR))

import make_cal_distribution as cal
import make_harald_260820_distribution as harald

REPORT_DATE = datetime.date.today().isoformat()
# 2026-08-21: was a hardcoded "2026-08-20"-dated filename -- only mattered
# when running this standalone (its own main(), guarded below), since the
# active GUI pipeline never calls it; but a stale hardcoded date silently
# targets a wrong/nonexistent file on any later day, same class of bug
# already fixed in update_harald_260820_summary_by_vendor.py's own
# SUMMARY_PATH. Kept in sync with that script's own REPORT_DATE convention.
SUMMARY_PATH = harald.PATHS.report_dir / f"{REPORT_DATE}_Harald_260820_Summary.xlsx"
def _resolve_nested_bench_dir() -> Path:
    """Whichever folder vendor/config-group derivation and hex-id -> serial
    mapping should read from.

    2026-08-21: this used to be hardcoded to _BASE_DIR/"Bench_40pcs"/
    "Harald" -- this workspace's own folder-naming convention, which a
    colleague's deployment didn't share at all (their Bench folder, picked
    via the GUI's own Folders card, was just a top-level "Bench/" with no
    "Bench_40pcs" wrapper) -- FileNotFoundError on their machine. Changed
    to always follow Harald_260820's own currently-configured bench_dir
    instead -- but that broke the OTHER everyday case: bench_dir pointed
    at the FLAT default pull (Bench_40pcs/Harald_260820, 40 files directly,
    no subfolders at all -- carries no vendor/serial info of its own, see
    this module's own docstring), which then found zero groups and crashed
    the exact same way (empty vendor_order -> a zero-width merge_cells).

    2026-08-25: prefer the active bench_dir only when IT is itself
    vendor/config-nested (has real subdirectories) -- that's what makes a
    100-DUT pull's DNNJ/DSKJ/SNKJ/SSNJ split work, since that's where the
    grouping actually lives. Otherwise (active bench_dir is flat) fall
    back to Bench_40pcs/Harald, the original always-nested (Skyworks/Sony)
    reference copy -- exactly the reasoning this constant existed for
    before the brief always-follow-bench_dir change above."""
    active = harald.PATHS.bench_dir
    if active.is_dir() and any(p.is_dir() for p in active.iterdir()):
        return active
    return _BASE_DIR / "Bench_40pcs" / "Harald"


NESTED_BENCH_DIR = _resolve_nested_bench_dir()


def build_hex_to_serial(nested_dir: Path) -> dict:
    mapping = {}
    for p in nested_dir.rglob("*.csv"):
        if p.name.startswith("."):  # macOS AppleDouble junk (._foo.csv)
            continue
        mapping[cal.dut_id_of(p)] = p.parent.name
    return mapping


def vendor_lists_from_nested_bench(nested_dir: Path, ate_path: Path | None = None) -> dict:
    """-> {GROUP_LABEL: [serial, ...]}, one entry per top-level subfolder
    of nested_dir, derived directly from its own {subfolder}/{DUT
    serial}/*.csv structure -- the subfolder's own name IS the label,
    verbatim (just uppercased), whatever it happens to be.

    2026-08-25: per reviewer request, this groups by whatever the real
    Bench folders actually are -- Skyworks/Sony for the original 40-DUT
    pull (2 groups), or DNNJ/DSKJ/SNKJ/SSNJ for a 100-DUT pull (4 groups,
    one per physical test config) -- rather than trying to collapse
    multiple config folders into a smaller number of "real vendors" (an
    earlier version of this function parsed a VENDOR(CONFIG) pattern out
    of the ATE filenames to do exactly that; reverted, since the reviewer
    wants the actual folder split preserved, not re-collapsed). `ate_path`
    is accepted but unused -- kept as an optional parameter only so
    existing callers passing it don't need to change.

    The DUT serial for each CSV is whatever its immediate parent folder is,
    found via rglob (not a fixed depth) so the one file that sits an extra
    level deep -- see this module's own docstring -- still lands under its
    correct serial instead of being missed by a depth-2-only glob."""
    groups: dict[str, list] = {}
    for subfolder in sorted(nested_dir.iterdir()):
        if not subfolder.is_dir():
            continue
        label = subfolder.name.upper()
        serials = set()
        for p in subfolder.rglob("*.csv"):
            if p.name.startswith("."):
                continue
            serials.add(p.parent.name)
        groups[label] = sorted(serials)
    return groups


def read_ate_per_dut(ate_csv: Path) -> dict:
    """-> {Test Item: {dut_serial: value}} -- same value-validity rule as
    cal._values_from_row (blank/NA skipped, |v|>=1e6 no-read sentinel
    dropped), just keeping the DUT-serial pairing cal.read_ate() itself
    discards."""
    idx = {}
    with ate_csv.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        header = next(reader)
        serials = header[3:]
        for row in reader:
            if not row:
                continue
            name = row[0].strip()
            rec = {}
            for serial, raw in zip(serials, row[3:]):
                v = raw.strip()
                if v in ("", "NA"):
                    continue
                try:
                    f = float(v)
                except ValueError:
                    continue
                if abs(f) >= 1e6:
                    continue
                rec[serial] = f
            if rec:
                idx[name] = rec
    return idx


def read_ate_per_dut_auto(path: Path) -> dict:
    """Same return shape as read_ate_per_dut() -- {Test Item: {dut_serial:
    value}} -- but accepts either a single ATE csv file or a FOLDER of them
    (see make_cal_distribution.read_ate_auto's own docstring for why: a
    100-DUT pull split ATE across 4 config-specific files). Each file's
    DUT-serial set is disjoint (confirmed 2026-08-21: 25 non-overlapping
    serials per file), so merging is a plain dict union per item -- no
    serial ever needs reconciling between files."""
    if not path.is_dir():
        return read_ate_per_dut(path)
    merged: dict[str, dict] = {}
    for f in sorted(p for p in path.glob("*.csv") if not p.name.startswith(".")):
        for name, rec in read_ate_per_dut(f).items():
            merged.setdefault(name, {}).update(rec)
    return merged


def bench_values_by_serial(rec: dict, hex_to_serial: dict) -> dict:
    out = {}
    for value, hexid in rec["values"]:
        serial = hex_to_serial.get(hexid)
        if serial:
            out[serial] = value
    return out


def main() -> None:
    print("Building hex-id -> DUT-serial map from the reference Bench copy...")
    hex_to_serial = build_hex_to_serial(NESTED_BENCH_DIR)
    print(f"  {len(hex_to_serial)} DUT(s) mapped")

    print("Reading Bench_Harald_260820 CSVs...")
    bench_items = cal.read_bench(harald.PATHS.bench_dir)
    print(f"  {len(bench_items)} test item(s) found")

    print(f"Parsing ATE log ({harald.ATE_PATH.name}) per-DUT...")
    ate_by_item = read_ate_per_dut(harald.ATE_PATH)
    ate_serials = set()
    for rec in ate_by_item.values():
        ate_serials.update(rec)
    all_serials = sorted(set(hex_to_serial.values()) | ate_serials)
    print(f"  {len(all_serials)} DUT serial(s) total")

    # Same qualifying set/order as build_harald_260820_xlsx_summary.py's own
    # jobs (>=1 Bench limit) -- reuses cal.build_jobs directly so this sheet
    # lists the identical Test Items as the existing summary sheet.
    ate_items_std = cal.read_ate(harald.ATE_PATH)
    jobs, _n_no_ate = cal.build_jobs(bench_items, ate_items_std)
    print(f"  {len(jobs)} qualifying item(s)")

    n_cols = 1 + 2 * len(all_serials)
    print("Building rows...")
    rows = []
    for job in jobs:
        raw_name = job["name"]
        item = harald.clean(raw_name)
        bench_by_serial = bench_values_by_serial(bench_items[raw_name], hex_to_serial)
        ate_by_serial = ate_by_item.get(raw_name, {})
        row = [item]
        for serial in all_serials:
            row.append(bench_by_serial.get(serial))
        for serial in all_serials:
            row.append(ate_by_serial.get(serial))
        rows.append(row)

    print("Writing sheet into existing workbook...")
    wb = openpyxl.load_workbook(SUMMARY_PATH)
    if "Raw Values" in wb.sheetnames:
        del wb["Raw Values"]
    ws = wb.create_sheet("Raw Values")

    ws.cell(row=1, column=1, value="Test Item")
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)

    col = 2
    for group_label, count in (("Bench", len(all_serials)), ("ATE", len(all_serials))):
        start = col
        for serial in all_serials:
            ws.cell(row=2, column=col, value=serial)
            col += 1
        ws.cell(row=1, column=start, value=group_label)
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=col - 1)

    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=2, max_col=n_cols):
        for cell in row:
            cell.font = header_font
            cell.alignment = center
    for i in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 12

    for r, row in enumerate(rows, start=3):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)

    ws.freeze_panes = "B3"
    wb.save(SUMMARY_PATH)
    print(f"\nSaved: {SUMMARY_PATH}  (added sheet 'Raw Values', {len(rows)} row(s) x {n_cols} col(s))")


if __name__ == "__main__":
    main()
