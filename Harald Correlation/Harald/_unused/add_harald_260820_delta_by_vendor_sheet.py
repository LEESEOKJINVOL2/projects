"""Add a "Delta by Vendor" sheet to the existing Harald_260820 xlsx summary
(report/Harald_260820/2026-08-20_Harald_260820_Summary.xlsx), per user
request 2026-08-20: one row per Test Item, columns grouped 3 levels deep --
Bench / ATE / Delta (top), then SKYWORKS / SONY (vendor), then individual
DUT serial ("2did") -- matching the header style shown in the user's
screenshot (Bench -> SKYWORKS -> DUT-serial columns), extended with the ATE
and Delta blocks alongside Bench since per-DUT delta needs both sides.

Delta here is per-DUT: |Bench_value - ATE_value| for THAT SAME DUT (not the
aggregate Mean-vs-Mean delta the first summary sheet already has) -- blank
if either side is missing for that DUT. Vendor grouping/DUT lists come from
the user-provided Downloads/"260818_2did 40pcs.txt" (SKYWORKS 20 + SONY 20),
parsed the same way as build_harald_260820_xlsx_summary_by_vendor.py.

Reuses add_harald_260820_raw_values_sheet.py's hex-id -> serial mapping and
per-DUT ATE parsing (imported, not modified) -- this file only adds the
new sheet's own layout/delta computation.
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

_THIS_DIR = Path(__file__).resolve().parent
_BASE_DIR = _THIS_DIR.parent
sys.path.insert(0, str(_BASE_DIR))
sys.path.insert(0, str(_BASE_DIR / "Cal"))
sys.path.insert(0, str(_THIS_DIR))

import xlsx_summary_writer as xw
import make_cal_distribution as cal
import make_harald_260820_distribution as harald
import add_harald_260820_raw_values_sheet as raw_sheet
import build_harald_260820_xlsx_summary_by_vendor as by_vendor

SUMMARY_PATH = harald.PATHS.report_dir / "2026-08-20_Harald_260820_Summary.xlsx"


def main() -> None:
    print(f"Parsing vendor DUT list ({by_vendor.VENDOR_LIST_PATH.name})...")
    vendor_lists = by_vendor.parse_vendor_lists(by_vendor.VENDOR_LIST_PATH)
    vendor_order = [v for v in ("SKYWORKS", "SONY") if v in vendor_lists]
    for v in vendor_order:
        vendor_lists[v] = sorted(vendor_lists[v])
        print(f"  {v}: {len(vendor_lists[v])} DUT(s)")

    print("Building hex-id -> DUT-serial map from the reference Bench copy...")
    hex_to_serial = raw_sheet.build_hex_to_serial(raw_sheet.NESTED_BENCH_DIR)
    print(f"  {len(hex_to_serial)} DUT(s) mapped")

    print("Reading Bench_Harald_260820 CSVs...")
    bench_items = cal.read_bench(harald.PATHS.bench_dir)
    print(f"  {len(bench_items)} test item(s) found")

    print(f"Parsing ATE log ({harald.ATE_PATH.name}) per-DUT...")
    ate_by_item = raw_sheet.read_ate_per_dut_auto(harald.ATE_PATH)  # ATE_PATH may be one file or a folder

    ate_items_std = cal.read_ate_auto(harald.ATE_PATH)  # ATE_PATH may be one file or a folder
    jobs, _n_no_ate = cal.build_jobs(bench_items, ate_items_std)
    print(f"  {len(jobs)} qualifying item(s)")

    print("Building rows...")
    rows = []
    for job in jobs:
        raw_name = job["name"]
        item = harald.clean(raw_name)
        bench_by_serial = raw_sheet.bench_values_by_serial(bench_items[raw_name], hex_to_serial)
        ate_by_serial = ate_by_item.get(raw_name, {})

        row = [item]
        for group in ("bench", "ate", "delta"):
            for vendor in vendor_order:
                for serial in vendor_lists[vendor]:
                    b = bench_by_serial.get(serial)
                    a = ate_by_serial.get(serial)
                    if group == "bench":
                        row.append(b)
                    elif group == "ate":
                        row.append(a)
                    else:
                        row.append(xw.delta(b, a))
        rows.append(row)

    print("Writing sheet into existing workbook...")
    wb = openpyxl.load_workbook(SUMMARY_PATH)
    sheet_name = "Delta by Vendor"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    ws.cell(row=1, column=1, value="Test Item")
    ws.merge_cells(start_row=1, start_column=1, end_row=3, end_column=1)

    col = 2
    for group_label in ("Bench", "ATE", "Delta"):
        group_start = col
        for vendor in vendor_order:
            vendor_start = col
            for serial in vendor_lists[vendor]:
                ws.cell(row=3, column=col, value=serial)
                col += 1
            ws.cell(row=2, column=vendor_start, value=vendor)
            ws.merge_cells(start_row=2, start_column=vendor_start, end_row=2, end_column=col - 1)
        ws.cell(row=1, column=group_start, value=group_label)
        ws.merge_cells(start_row=1, start_column=group_start, end_row=1, end_column=col - 1)
    n_cols = col - 1

    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=3, max_col=n_cols):
        for cell in row:
            cell.font = header_font
            cell.alignment = center
    for i in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 12

    for r, row in enumerate(rows, start=4):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)

    ws.freeze_panes = "B4"
    wb.save(SUMMARY_PATH)
    print(f"\nSaved: {SUMMARY_PATH}  (added sheet '{sheet_name}', {len(rows)} row(s) x {n_cols} col(s))")


if __name__ == "__main__":
    main()
