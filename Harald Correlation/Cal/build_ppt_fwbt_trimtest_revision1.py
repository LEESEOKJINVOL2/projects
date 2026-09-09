"""Build the REAL FWBT_TRIMTEST PPT deck (Bench + ATE data, real plots) --
graduates the layout-only draft (build_ppt_fwbt_trimtest_draft.py, approved
2026-08-15) into a full report, per user request the same day.

Slide layout is Cal's own interleaved-table + stacked-image convention
(Cal/build_ppt_cal_distribution.py: ONE table -- item/band header, Bench/
ATE sub-header, stat rows -- then Bench image row on top, ATE image row (or
a "no ATE match" hole) below), extended with two explicit USL/LSL rows
(red text, merged across each band's Bench+ATE sub-columns) plus a "span"
ruler (axis vmax/vmin shown once per Bench/ATE row) matching the existing
BT_TX/build_ppt_pwr_evm_revision2.py convention the user pointed to as the
visual reference (its _add_no_ate_box / USL_COLOR / ATE_TEXT_COLOR /
_add_label / _add_ruler are all reused here verbatim, not modified in that
file). The ruler only works because every chart in a group now shares ONE
Y-axis range (make_bt_tx_trimtest_distribution.group_value_range), not a
per-item auto range. Grouping
(2G-left/5G-right by band) and slide titles reuse the already-approved
draft's clean_title/base_key/BAND_ORDER, and job data comes from
make_bt_tx_trimtest_distribution.py (Bench/ATE values + USL/LSL spec +
the PNGs that script already drew).
"""
from __future__ import annotations

import datetime
import statistics
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

_THIS_DIR = Path(__file__).resolve().parent  # Cal/ itself (2026-08-19: moved
# in from BT_TX_TRIMTEST/ per user request, see make_bt_tx_trimtest_
# distribution.py's matching note)
sys.path.insert(0, str(_THIS_DIR))

import make_cal_distribution as cal  # safe_filename (reused, not modified)
import make_bt_tx_trimtest_distribution as data

PATHS = data.PATHS
TEMPLATE_PATH = PATHS.ppt_template

SLIDE_WIDTH_IN = data.SLIDE_WIDTH_IN
LEFT_MARGIN = data.LEFT_MARGIN
RIGHT_MARGIN = data.RIGHT_MARGIN
CELL_GAP_H = data.CELL_GAP_H
LABEL_WIDTH = data.LABEL_WIDTH
TABLE_TOP = data.TABLE_TOP
HEADER_ROW_H = data.HEADER_ROW_H
TABLE_ROW_H = data.TABLE_ROW_H
TABLE_PLOT_GAP = data.TABLE_PLOT_GAP
IMAGE_GAP = data.IMAGE_GAP
CONTENT_AREA_TOP = data.CONTENT_AREA_TOP
CONTENT_AREA_BOTTOM = data.CONTENT_AREA_BOTTOM

# -- reused verbatim from BT_TX/build_ppt_pwr_evm_revision2.py's own red
# spec/ATE-hole convention, per user instruction to match that visual style --
ATE_BORDER_COLOR = RGBColor(0xC0, 0x00, 0x00)
ATE_TEXT_COLOR = RGBColor(0xC0, 0x00, 0x00)
USL_COLOR = RGBColor(0xB0, 0x00, 0x00)

REPORT_DATE = datetime.date.today().isoformat()
TITLE_TEXT = f"{PATHS.title_prefix} {PATHS.project} {PATHS.title_suffix}"
SUBTITLE_LINES = ["RF LAB1 Team", REPORT_DATE]


def _find_layout(prs, prefix, fallback):
    for i, layout in enumerate(prs.slide_layouts):
        if layout.name.startswith(prefix):
            return i
    return fallback


def _get_ph(slide, idx):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None


def _remove_placeholder(slide, idx: int) -> None:
    ph = _get_ph(slide, idx)
    if ph is not None:
        ph._element.getparent().remove(ph._element)


def _set_cell(cell, text, size=7, bold=True, color=None):
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Pt(2)
    cell.margin_right = Pt(2)
    cell.margin_top = Pt(1)
    cell.margin_bottom = Pt(1)
    tf = cell.text_frame
    tf.word_wrap = True
    lines = (text if text else " ").split("\n")
    tf.text = lines[0]
    for line in lines[1:]:
        tf.add_paragraph().text = line
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        run = p.runs[0] if p.runs else p.add_run()
        run.font.size = Pt(size)
        run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color


def _add_label(slide, left, top, width, height, text, size=9, bold=True):
    """Reused (reimplemented, not imported) from BT_TX/build_ppt_pwr_evm_
    revision2.py's own _add_label -- a plain borderless textbox."""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    return box


def _add_ruler(slide, left, top, width, height, vmin, vmax, row_label=None):
    """"Span" ruler -- vmax flush with the image row's top edge, vmin flush
    with its bottom edge. Per user request 2026-08-15 (pointing to the
    yellow-highlighted "36"/"0" labels on an existing BT_TX_DEVM Peak
    slide) -- reused verbatim from BT_TX/build_ppt_pwr_evm_revision2.py's
    own _add_ruler. Only meaningful because every chart drawn into this
    group's image row already shares the SAME vmin/vmax (see
    make_bt_tx_trimtest_distribution.group_value_range) -- a per-column
    ruler would be needed instead if that weren't true.

    row_label (2026-08-15, later same day): "Bench"/"ATE" text centered
    between the vmax/vmin numbers, identifying which stacked image row this
    ruler belongs to -- per user screenshot circling where that label
    should go on an existing slide."""
    tick_h = 0.2
    _add_label(slide, left, top, width, tick_h, f"{vmax:.0f}", size=9, bold=True)
    _add_label(slide, left, top + height - tick_h, width, tick_h, f"{vmin:.0f}", size=9, bold=True)
    if row_label:
        label_h = 0.4
        label_top = top + (height - label_h) / 2
        _add_label(slide, left, label_top, width, label_h, row_label, size=16, bold=True)


def _add_no_ate_box(slide, left, top, width, height):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    box.fill.background()
    box.line.color.rgb = ATE_BORDER_COLOR
    box.line.width = Pt(1)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "No ATE match"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = ATE_TEXT_COLOR


def _stats(values):
    if not values:
        return None
    mean = statistics.mean(values)
    std = statistics.stdev(values) if len(values) > 1 else 0.0
    return mean, min(values), max(values), std


def _fmt_stat(values):
    s = _stats(values)
    if s is None:
        return "", ""
    mean, vmin, vmax, std = s
    return f"{mean:.4g}±{std:.4g}", f"{vmin:.4g} / {vmax:.4g}"


def _fmt_limit(v):
    return f"{v:g}" if v is not None else ""


def _add_interleaved_table(slide, left, top, table_width, jobs):
    n_items = len(jobs)
    n_data_cols = 2 * n_items
    n_cols = 1 + n_data_cols
    n_rows = 6  # band header + Bench/ATE subheader + Mean+-Std + Min/Max + USL + LSL

    grid_width = table_width - LABEL_WIDTH
    cell_width = (grid_width - (n_data_cols - 1) * CELL_GAP_H) / n_data_cols
    col_widths = [LABEL_WIDTH] + [cell_width] * n_data_cols
    row_heights = [HEADER_ROW_H] + [TABLE_ROW_H] * (n_rows - 1)
    total_width = sum(col_widths)
    total_height = sum(row_heights)

    gframe = slide.shapes.add_table(n_rows, n_cols, Inches(left), Inches(top),
                                     Inches(total_width), Inches(total_height))
    table = gframe.table
    table.first_row = False
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)
    for r, h in enumerate(row_heights):
        table.rows[r].height = Inches(h)

    _set_cell(table.cell(0, 0), "Band", size=7)
    _set_cell(table.cell(1, 0), "", size=7)
    _set_cell(table.cell(2, 0), "Mean±Std", size=7)
    _set_cell(table.cell(3, 0), "Min/Max", size=7)
    _set_cell(table.cell(4, 0), "USL", size=7, color=USL_COLOR)
    _set_cell(table.cell(5, 0), "LSL", size=7, color=USL_COLOR)

    for ci, job in enumerate(jobs):
        c1 = 1 + ci * 2
        merged = table.cell(0, c1)
        merged.merge(table.cell(0, c1 + 1))
        _set_cell(merged, job["col_label"], size=7)
        _set_cell(table.cell(1, c1), "Bench", size=7)
        _set_cell(table.cell(1, c1 + 1), "ATE", size=7)

        b_mean, b_minmax = _fmt_stat(job["bench_values"])
        a_mean, a_minmax = _fmt_stat(job["ate_values"] or [])
        _set_cell(table.cell(2, c1), b_mean, size=6, bold=True)
        _set_cell(table.cell(2, c1 + 1), a_mean, size=6, bold=True)
        _set_cell(table.cell(3, c1), b_minmax, size=6, bold=True)
        _set_cell(table.cell(3, c1 + 1), a_minmax, size=6, bold=True)

        usl_cell = table.cell(4, c1)
        usl_cell.merge(table.cell(4, c1 + 1))
        _set_cell(usl_cell, _fmt_limit(job["usl"]), size=6, bold=True, color=USL_COLOR)
        lsl_cell = table.cell(5, c1)
        lsl_cell.merge(table.cell(5, c1 + 1))
        _set_cell(lsl_cell, _fmt_limit(job["lsl"]), size=6, bold=True, color=USL_COLOR)

    return total_height


def _add_image_stretched(slide, path: Path, left, top, width, height):
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))


def add_item_slide(prs, layout_idx, title_text, jobs):
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    ph_title = _get_ph(slide, 1) or _get_ph(slide, 0)
    if ph_title:
        ph_title.text = title_text
    _remove_placeholder(slide, 12)

    usable_left = LEFT_MARGIN
    usable_right = SLIDE_WIDTH_IN - RIGHT_MARGIN
    n_cols = len(jobs)
    cell_width = (usable_right - usable_left - LABEL_WIDTH - (n_cols - 1) * CELL_GAP_H) / n_cols

    table_h = _add_interleaved_table(slide, usable_left, TABLE_TOP, usable_right - usable_left, jobs)

    plots_top = TABLE_TOP + table_h + TABLE_PLOT_GAP
    plots_bottom = CONTENT_AREA_BOTTOM
    image_h = (plots_bottom - plots_top - IMAGE_GAP) / 2
    bench_top = plots_top
    ate_top = bench_top + image_h + IMAGE_GAP

    vmin, vmax = jobs[0]["value_range"]  # shared across the whole group
    _add_ruler(slide, usable_left, bench_top, LABEL_WIDTH, image_h, vmin, vmax, row_label="Bench")
    _add_ruler(slide, usable_left, ate_top, LABEL_WIDTH, image_h, vmin, vmax, row_label="ATE")

    for ci, job in enumerate(jobs):
        cell_left = usable_left + LABEL_WIDTH + ci * (cell_width + CELL_GAP_H)
        if job["bench_png"].exists():
            _add_image_stretched(slide, job["bench_png"], cell_left, bench_top, cell_width, image_h)
        if job["ate_values"] is not None and job["ate_png"].exists():
            _add_image_stretched(slide, job["ate_png"], cell_left, ate_top, cell_width, image_h)
        else:
            _add_no_ate_box(slide, cell_left, ate_top, cell_width, image_h)

    return slide


def add_title_slide(prs):
    layout_idx = _find_layout(prs, "2_", 0)
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    ph_title = _get_ph(slide, 0)
    if ph_title:
        ph_title.text = TITLE_TEXT
    ph_subtitle = _get_ph(slide, 1)
    if ph_subtitle:
        tf = ph_subtitle.text_frame
        tf.text = SUBTITLE_LINES[0]
        for line in SUBTITLE_LINES[1:]:
            p = tf.add_paragraph()
            p.text = line
    return slide


def _reorder_slides(prs, ordered_slides) -> None:
    xml_slides = prs.slides._sldIdLst
    id_by_rid = {e.get(qn("r:id")): e for e in list(xml_slides)}
    part_to_rid = {prs.part.related_part(rid): rid for rid in id_by_rid}
    for e in list(xml_slides):
        xml_slides.remove(e)
    for slide in ordered_slides:
        rid = part_to_rid[slide.part]
        xml_slides.append(id_by_rid[rid])


def main() -> None:
    prs = Presentation(TEMPLATE_PATH)
    content_layout_idx = _find_layout(prs, "1_", 5)

    print("Loading Test Item spec + Bench/ATE data...")
    specs = data.load_item_specs()
    bench_items = cal.read_bench(PATHS.bench_dir)
    ate_items = cal.read_ate(PATHS.ate_log_path)
    jobs = data.build_jobs(specs, bench_items, ate_items)
    groups = data.build_groups(jobs)
    print(f"{len(groups)} group(s) (slides)")

    bench_dir_out = data.bench_png_dir()
    ate_dir_out = data.ate_png_dir()
    for g in groups.values():
        for job in g["jobs"]:
            fname = cal.safe_filename(job["cleaned_name"])
            job["bench_png"] = bench_dir_out / f"{fname}.png"
            job["ate_png"] = ate_dir_out / f"{fname}.png"

    ordered = []
    title_slide = add_title_slide(prs)
    ordered.append(title_slide)

    n_missing_png = 0
    n_groups = len(groups)
    # 2026-08-25: same reasoning as Harald/build_ppt_harald_trimtest_
    # revision1.py's own note -- ReportBench's GUI parses "[n/total]"-
    # shaped lines to drive this step's progress bar, and this loop used to
    # print nothing between "N group(s)" and "Total content slides",
    # leaving that bar empty for however long a large deck takes.
    for n, key in enumerate(sorted(groups), 1):
        g = groups[key]
        for job in g["jobs"]:
            if not job["bench_png"].exists():
                n_missing_png += 1
        slide = add_item_slide(prs, content_layout_idx, g["title"], g["jobs"])
        ordered.append(slide)
        if n % 200 == 0 or n == n_groups:
            print(f"  [{n}/{n_groups}] {g['title']}")

    if n_missing_png:
        print(f"WARNING: {n_missing_png} Bench PNG(s) missing on disk -- "
              f"run make_bt_tx_trimtest_distribution.py first")

    print(f"Total content slides: {len(ordered) - 1}")
    _reorder_slides(prs, ordered)

    PATHS.report_dir.mkdir(parents=True, exist_ok=True)
    out_path = PATHS.report_dir / f"{REPORT_DATE}_FWBT_TRIMTEST_Revision1.pptx"
    prs.save(out_path)
    print(f"\nSaved: {out_path}")
    print(f"Total slides: {len(prs.slides._sldIdLst)}")


if __name__ == "__main__":
    main()
