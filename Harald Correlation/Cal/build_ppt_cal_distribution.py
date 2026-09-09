"""Build the Cal PPT deck: one slide per (PATNAME section, page) chunk of
qualifying items.

Cal has no packet_type/band/channel-style dimension the way BT_TX/BT_RX do --
its ~3847 qualifying items (>=1 Bench limit, see make_cal_distribution.py)
are grouped only by the Bench log's own "PATNAME:" section header (e.g.
"CLKMEASX_XXXXXXXX", "BT2GTXLD_OTRIMXXX"), 1043 such sections, averaging
~3.7 qualifying items each but ranging up to 601 in the largest single
section. COLS_PER_SLIDE paginates any section larger than that many items
across multiple slides (title gets a "(page/total)" suffix) rather than
shrinking columns to fit -- an item name here is a full ~40-60 char string
(unlike BT_TX's short "1DH5_2402MHz..." combo labels), so column width is
kept wide enough for a 3-4 line wrapped header to stay legible instead of
maximizing columns-per-slide.

Layout per slide (same interleaved-table + stacked-image convention as
BT_RX's build_ppt_bt_rx_revision1.py, reused near-verbatim -- see that
module's docstring for why there's no left-side ruler here either: each
item's Bench/ATE pair auto-ranges to its own data/spec, not a slide-wide
shared axis):
  * ONE table: item name (wrapped) header, Bench/ATE sub-header, Spec
    (LSL/USL), Mean+/-Std, Min/Max.
  * Two stacked image rows below (Bench on top, ATE below, ATE cell left
    blank via the same `if ate_path.exists()` guard when extract_mode
    forces ATE absent or a given item had no ATE match).

No section-divider slides -- 1043 sections is far too many for a divider
each, and the interleaved table's own header already names the section via
the slide title ("Cal_{section}" [+ page suffix]).
"""
from __future__ import annotations

import datetime
import statistics
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

import make_cal_distribution as cal

PATHS = cal.PATHS
TEMPLATE_PATH = PATHS.ppt_template

# COLS_PER_SLIDE lives on make_cal_distribution (cal.COLS_PER_SLIDE) -- it's
# the single source of truth both scripts read, since make_cal_distribution
# also needs it (via build_pages) to size each PNG's own figsize correctly.

SLIDE_WIDTH_IN = 13.338542213473316
LEFT_MARGIN = 0.05
RIGHT_MARGIN = 0.05
CELL_GAP_H = 0.03
LABEL_WIDTH = 0.9

TABLE_TOP = 1.05
TABLE_ROW_H = 0.2
# 3-4 line wrapped item-name header needs more room than one plain data row
# -- same HEADER_ROW_H-vs-TABLE_ROW_H split BT_RX uses, see its docstring.
HEADER_ROW_H = 0.55
TABLE_PLOT_GAP = 0.35
IMAGE_GAP = 0.15
CONTENT_AREA_TOP = 1.5908508311461067
CONTENT_AREA_BOTTOM = CONTENT_AREA_TOP + 5.240927384076991

REPORT_DATE = datetime.date.today().isoformat()
TITLE_TEXT = f"{PATHS.title_prefix} {PATHS.project} {PATHS.title_suffix}"
SUBTITLE_LINES = ["RF LAB1 Team", REPORT_DATE]


def _find_layout(prs, prefix, fallback):
    for i, layout in enumerate(prs.slide_layouts):
        if layout.name.startswith(prefix):
            return i
    return fallback


def _get_ph(slide, idx):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None


def _remove_placeholder(slide, idx: int) -> None:
    ph = _get_ph(slide, idx)
    if ph is not None:
        ph._element.getparent().remove(ph._element)


def _set_cell(cell, text, size=8, bold=True, color=None):
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Pt(2)
    cell.margin_right = Pt(2)
    cell.margin_top = Pt(1)
    cell.margin_bottom = Pt(1)
    tf = cell.text_frame
    tf.word_wrap = True
    tf.text = text if text else " "
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        run = p.runs[0] if p.runs else p.add_run()
        run.font.size = Pt(size)
        run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color


def _stats(values):
    if not values:
        return None
    return (statistics.mean(values), min(values), max(values),
            statistics.stdev(values) if len(values) > 1 else 0.0)


def _fmt_stat(values):
    s = _stats(values)
    if s is None:
        return "", ""
    mean, vmin, vmax, std = s
    return f"{mean:.4g}±{std:.4g}", f"{vmin:.4g} / {vmax:.4g}"


def _fmt_spec(llim, ulim):
    lo = f"{llim:.4g}" if llim is not None else "--"
    hi = f"{ulim:.4g}" if ulim is not None else "--"
    return f"{lo} ~ {hi}"


def wrap_item_name(name: str, max_len: int = 16) -> str:
    return cal.wrap_label(name, max_len=max_len)


def _add_interleaved_table(slide, left, top, table_width, jobs):
    n_items = len(jobs)
    n_data_cols = 2 * n_items
    n_cols = 1 + n_data_cols
    n_rows = 5  # item header + Bench/ATE subheader + Spec + Mean+-Std + Min/Max

    grid_width = table_width - LABEL_WIDTH
    cell_width = (grid_width - (n_data_cols - 1) * CELL_GAP_H) / n_data_cols
    col_widths = [LABEL_WIDTH] + [cell_width] * n_data_cols
    total_width = sum(col_widths)
    row_heights = [HEADER_ROW_H] + [TABLE_ROW_H] * (n_rows - 1)
    total_height = sum(row_heights)

    gframe = slide.shapes.add_table(n_rows, n_cols, Inches(left), Inches(top),
                                     Inches(total_width), Inches(total_height))
    table = gframe.table
    table.first_row = False
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)
    for r, h in enumerate(row_heights):
        table.rows[r].height = Inches(h)

    _set_cell(table.cell(0, 0), "Item", size=7)
    _set_cell(table.cell(1, 0), "", size=7)
    _set_cell(table.cell(2, 0), "Spec (LSL/USL)", size=7)
    _set_cell(table.cell(3, 0), "Mean±Std", size=7)
    _set_cell(table.cell(4, 0), "Min/Max", size=7)

    for ci, job in enumerate(jobs):
        c1 = 1 + ci * 2
        merged = table.cell(0, c1)
        merged.merge(table.cell(0, c1 + 1))
        _set_cell(merged, wrap_item_name(job["name"]), size=6)
        _set_cell(table.cell(1, c1), "Bench", size=7)
        _set_cell(table.cell(1, c1 + 1), "ATE", size=7)

        llim, ulim = cal.effective_spec(job)
        _set_cell(table.cell(2, c1), _fmt_spec(job["bench_llim"], job["bench_ulim"]), size=6)
        _set_cell(table.cell(2, c1 + 1), _fmt_spec(llim, ulim), size=6)

        b_mean, b_minmax = _fmt_stat(job["bench_values"])
        a_mean, a_minmax = _fmt_stat(job["ate_values"])
        _set_cell(table.cell(3, c1), b_mean, size=7, bold=True)
        _set_cell(table.cell(3, c1 + 1), a_mean, size=7, bold=True)
        _set_cell(table.cell(4, c1), b_minmax, size=7, bold=True)
        _set_cell(table.cell(4, c1 + 1), a_minmax, size=7, bold=True)

    return total_height


def _add_image_stretched(slide, path: Path, left, top, width, height):
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))


def add_item_slide(prs, layout_idx, title_text, jobs, bench_dir, ate_dir):
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    ph_title = _get_ph(slide, 1) or _get_ph(slide, 0)
    if ph_title:
        ph_title.text = title_text
    _remove_placeholder(slide, 12)

    usable_left = LEFT_MARGIN
    usable_right = SLIDE_WIDTH_IN - RIGHT_MARGIN
    n_cols = len(jobs)
    cell_width = (usable_right - usable_left - LABEL_WIDTH - (n_cols - 1) * CELL_GAP_H) / n_cols

    table_h = _add_interleaved_table(slide, usable_left, TABLE_TOP, usable_right - usable_left, jobs)

    plots_top = TABLE_TOP + table_h + TABLE_PLOT_GAP
    plots_bottom = CONTENT_AREA_BOTTOM
    image_h = (plots_bottom - plots_top - IMAGE_GAP) / 2
    bench_top = plots_top
    ate_top = bench_top + image_h + IMAGE_GAP

    for ci, job in enumerate(jobs):
        cell_left = usable_left + LABEL_WIDTH + ci * (cell_width + CELL_GAP_H)
        fname = cal.safe_filename(job["name"]) + ".png"
        bench_path = bench_dir / fname
        ate_path = ate_dir / fname
        if bench_path.exists():
            _add_image_stretched(slide, bench_path, cell_left, bench_top, cell_width, image_h)
        if ate_path.exists():
            _add_image_stretched(slide, ate_path, cell_left, ate_top, cell_width, image_h)

    return slide


def add_title_slide(prs):
    layout_idx = _find_layout(prs, "2_", 0)
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    ph_title = _get_ph(slide, 0)
    if ph_title:
        ph_title.text = TITLE_TEXT
    ph_subtitle = _get_ph(slide, 1)
    if ph_subtitle:
        tf = ph_subtitle.text_frame
        tf.text = SUBTITLE_LINES[0]
        for line in SUBTITLE_LINES[1:]:
            p = tf.add_paragraph()
            p.text = line
    return slide


def _reorder_slides(prs, ordered_slides) -> None:
    xml_slides = prs.slides._sldIdLst
    id_by_rid = {e.get(qn("r:id")): e for e in list(xml_slides)}
    part_to_rid = {prs.part.related_part(rid): rid for rid in id_by_rid}
    for e in list(xml_slides):
        xml_slides.remove(e)
    for slide in ordered_slides:
        rid = part_to_rid[slide.part]
        xml_slides.append(id_by_rid[rid])


def build_all_jobs():
    bench_items = cal.read_bench(PATHS.bench_dir)
    ate_items = cal.read_ate(PATHS.ate_log_path)
    jobs, n_no_ate = cal.build_jobs(bench_items, ate_items)
    return jobs, n_no_ate, len(bench_items)


def main() -> None:
    prs = Presentation(TEMPLATE_PATH)
    content_layout_idx = _find_layout(prs, "1_", 5)

    print("Reading Bench_Cal CSVs + ATE log...")
    jobs, n_no_ate, n_bench_items = build_all_jobs()
    print(f"=== items with >=1 Bench limit: {len(jobs)}/{n_bench_items} "
          f"(ATE match: {len(jobs) - n_no_ate}, no ATE match: {n_no_ate}) ===")

    # SAME pagination make_cal_distribution.py used to size each PNG's own
    # figsize (see cal.build_pages' docstring) -- calling that one shared
    # function here instead of re-deriving the grouping guarantees this
    # deck's slide layout always matches what each PNG was actually drawn
    # for, so this doesn't drift out of sync with the PNG generation step.
    pages = cal.build_pages(jobs)
    n_sections = len({p["section"] for p in pages})
    print(f"{n_sections} PATNAME section(s) with a qualifying item")

    bench_dir = cal.bench_png_dir()
    ate_dir = cal.ate_png_dir() if PATHS.extract_mode != "bench" else Path("__no_such_dir__")

    ordered = []
    for page in pages:
        title = f"Cal_{page['section']}"
        if page["n_pages"] > 1:
            title += f" ({page['page']}/{page['n_pages']})"
        slide = add_item_slide(prs, content_layout_idx, title, page["jobs"], bench_dir, ate_dir)
        ordered.append(slide)

    print(f"{len(pages)} content slide(s) ({cal.COLS_PER_SLIDE} item(s)/slide max)")

    title_slide = add_title_slide(prs)
    ordered = [title_slide, *ordered]
    _reorder_slides(prs, ordered)

    report_dir = PATHS.report_dir
    report_dir.mkdir(parents=True, exist_ok=True)
    out_path = report_dir / f"{REPORT_DATE}_Cal_Distribution.pptx"
    prs.save(out_path)
    print(f"\nSaved: {out_path}")
    print(f"Total slides: {len(prs.slides._sldIdLst)}")


if __name__ == "__main__":
    main()
