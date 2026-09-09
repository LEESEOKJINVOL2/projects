"""Build BT_TX's single combined deliverable deck in ONE pass:
Power -> DEVM RMS/Peak/99pct -> ACP -> FreqAcc's sub-metrics, one title slide,
saved as {date}_BT_TX_Full_Revision2.pptx.

2026-07-28: replaces the old build-two-decks-then-merge pipeline. Previously
run_all.py ran three steps -- build_ppt_full_revision2.py (Power/DEVM/ACP),
build_ppt_freqacc_revision2.py (FreqAcc), then merge_ppt_full_revision2.py,
which stitched the two finished .pptx files together through PowerPoint COM's
Slides.InsertFromFile and then re-saved the result through python-pptx purely
to re-compress the zip (COM's own SaveAs barely deflates it). Measured on this
workspace those three steps were 8:56 + 1:12 + 9:37 = 19:45 of a 38:22 BT_TX
run, and ~9:37 of that was the merge alone.

Nothing about the merge was ever load-bearing: both source decks were built
off the same PPT-Template.pptx by the same slide-building functions in
build_ppt_pwr_evm_revision2.py, so appending FreqAcc's sections into the SAME
python-pptx Presentation produces the same deck without the round trip. What
that removes:
  * PowerPoint COM from the build path entirely (still used for visual
    verification -- see md/CLAUDE.md)
  * the taskkill of POWERPNT.EXE, which killed any unrelated presentation the
    user happened to have open
  * the re-compression pass, and with it the read-and-overwrite-the-same-file
    step that could destroy a ~30-minute deliverable if it failed partway
  * two intermediate .pptx files (~500MB + ~230MB) written and immediately
    superseded

Deliberately NOT deleted: build_ppt_full_revision2.py and
build_ppt_freqacc_revision2.py each still have a working main() that produces
its own standalone half-deck. They are simply no longer on run_all.py's path,
so the two intermediate files stop being produced by a normal pipeline run.
Run either script directly if you want one.

Peak memory is LOWER than the pipeline this replaces, not higher: one
Presentation holding all ~530 slides' images is roughly the size of the final
deck (~726MB, essentially the PNG bytes -- already-compressed PNGs get no zip
gain), whereas the merge step used to load the full ~1.1GB COM output into
python-pptx on top of having just written it.
"""

from __future__ import annotations

import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.dml.color import RGBColor

import build_ppt_pwr_evm_revision2 as PE
import build_ppt_full_revision2 as PWR_DEVM_ACP
import build_ppt_freqacc_revision2 as FREQACC

import sys
from pathlib import Path as _Path
sys.path.insert(0, str(_Path(__file__).resolve().parent.parent))
import paths as bt_tx_paths

PATHS = bt_tx_paths.get_paths("BT_TX")
REPORT_DATE = datetime.date.today().isoformat()

TOC_LINK_COLOR = RGBColor(0x1F, 0x4E, 0x79)  # dark blue, reads as a link


def _divider_label(divider) -> str:
    """The divider slide's own section-name text (see PE.make_divider_
    factory/_set_section_text) -- read back rather than threaded through
    separately, so the TOC can never drift from what's actually printed
    on each divider."""
    ph = PE._get_ph(divider, 1) or PE._get_ph(divider, 0)
    return ph.text_frame.text if ph else "?"


TOC_BLOCK_WIDTH_IN = 9.0  # centered column, NOT the full usable_left..usable_right
# chart-grid width -- that made row 1 (2026-08-15 v1) pin the label to the
# slide's left edge and the page range to its right edge with a huge dead
# gap between, which read as "spread across both sides" rather than a
# normal TOC list (2026-08-15 customer feedback, v2).
TOC_RULE_COLOR = RGBColor(0xD9, 0xD9, 0xD9)  # light gray row divider


def _add_toc_slide(prs, layout_idx, section_toc):
    """One Table-of-Contents slide, one row per section: "{label}" left,
    "{start}-{end}" page range right, BOTH clickable (2026-08-15 customer
    request: TOC page + hyperlink navigation). Each row is its own
    textbox (not a table cell) so python-pptx's shape-level
    click_action.target_slide can jump straight to that section's own
    divider slide -- reordering slides afterward doesn't break this,
    since the link targets the slide PART, not its position.

    Laid out in a centered TOC_BLOCK_WIDTH_IN-wide column (not the full
    chart-grid usable width) with a thin rule under each row, so label and
    page range sit close together like a normal table of contents instead
    of pinned to opposite slide edges."""
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    PE._remove_placeholder(slide, 12)
    ph_title = PE._get_ph(slide, 1) or PE._get_ph(slide, 0)
    if ph_title:
        ph_title.text = "Table of Contents"

    usable_left = PE.LEFT_MARGIN
    usable_right = PE.SLIDE_WIDTH_IN - PE.RIGHT_MARGIN
    block_width = min(TOC_BLOCK_WIDTH_IN, usable_right - usable_left)
    block_left = usable_left + (usable_right - usable_left - block_width) / 2

    top = PE.CONTENT_AREA_TOP
    available_h = PE.CONTENT_AREA_BOTTOM - top
    row_h = min(0.42, available_h / max(1, len(section_toc)))
    label_w = block_width * 0.7
    page_w = block_width * 0.3

    for i, (label, start, end, divider) in enumerate(section_toc):
        row_top = top + i * row_h

        label_box = slide.shapes.add_textbox(Inches(block_left), Inches(row_top), Inches(label_w), Inches(row_h))
        tf = label_box.text_frame
        tf.margin_left = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        run = tf.paragraphs[0].add_run()
        run.text = label
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.underline = True
        run.font.color.rgb = TOC_LINK_COLOR

        page_box = slide.shapes.add_textbox(Inches(block_left + label_w), Inches(row_top),
                                             Inches(page_w), Inches(row_h))
        tf2 = page_box.text_frame
        tf2.margin_left = 0
        tf2.margin_top = 0
        tf2.margin_bottom = 0
        tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.RIGHT
        run2 = p2.add_run()
        run2.text = f"{start}-{end}" if end != start else f"{start}"
        run2.font.size = Pt(14)
        run2.font.bold = True
        run2.font.color.rgb = TOC_LINK_COLOR

        label_box.click_action.target_slide = divider
        page_box.click_action.target_slide = divider

        rule = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(block_left), Inches(row_top + row_h - 0.02),
                                           Inches(block_left + block_width), Inches(row_top + row_h - 0.02))
        rule.line.color.rgb = TOC_RULE_COLOR
        rule.line.width = Pt(0.75)

    return slide


def main() -> None:
    prs = Presentation(PE.TEMPLATE_PATH)
    content_layout_idx = PE._find_layout(prs, "1_", 5)

    # ONE factory across both halves: Power consumes the template's slide 0 and
    # every later section (DEVM x3, ACP, then each FreqAcc sub-metric)
    # duplicates it. Handing the same factory to both build_sections() calls is
    # what keeps that invariant intact across the seam.
    new_divider = PE.make_divider_factory(prs)

    print("=== Power / DEVM / ACP sections ===")
    sections = PWR_DEVM_ACP.build_sections(prs, content_layout_idx, new_divider)

    print("\n=== FreqAcc sections ===")
    sections += FREQACC.build_sections(prs, content_layout_idx, new_divider)

    # Same config.xml-driven title as every other project's deck (see
    # paths.py's title_prefix/title_suffix docstrings). One title slide for the
    # combined deck -- the merge it replaces left the base deck's title slide
    # in place and appended FreqAcc's slides after it, so this matches.
    PE.TITLE_TEXT = f"{PATHS.title_prefix} {PATHS.project} {PATHS.title_suffix}"
    PE.SUBTITLE_LINES = ["RF LAB1 Team", REPORT_DATE]
    title_slide = PE.add_title_slide(prs)

    # 2026-08-15 (customer request): Table of Contents, one hyperlinked row
    # per section with its page range. 2026-08-27 (customer request): TOC
    # moved to page 2 (was page 1) -- title=page 1, TOC=page 2, each
    # section's own divider slide IS that section's first page (so
    # "Power: 3-22" means page 3 is the Power divider itself, pages 4-22
    # are its packet slides) -- unaffected by the title/TOC swap since
    # content still starts at page 3 either way.
    print("\nComputing Table of Contents page ranges...")
    section_toc = []
    page = 3
    for divider, slides in sections:
        label = _divider_label(divider)
        start = page
        end = page + len(slides)
        section_toc.append((label, start, end, divider))
        page = end + 1
    toc_slide = _add_toc_slide(prs, content_layout_idx, section_toc)
    for label, start, end, _divider in section_toc:
        print(f"  {label}: {start}-{end}")

    ordered = [title_slide, toc_slide]
    for divider, slides in sections:
        ordered.extend([divider, *slides])
    PE._reorder_slides(prs, ordered)

    report_dir = PATHS.report_dir
    report_dir.mkdir(parents=True, exist_ok=True)
    out_path = report_dir / f"{REPORT_DATE}_BT_TX_Full_Revision2.pptx"
    prs.save(out_path)
    print(f"\nSaved: {out_path}  ({out_path.stat().st_size / 1e6:.0f}MB)")
    print(f"Total slides: {len(prs.slides._sldIdLst)}  "
          f"({len(sections)} sections + 1 title + 1 TOC)")


if __name__ == "__main__":
    main()
