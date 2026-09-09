"""Give BT_TX the per-config/vendor summary feature Harald already has
(Harald/update_harald_260820_summary_by_vendor.py, the primary template this
mirrors closely), per user request 2026-08-25:

1. Add one "<Config> Summary" sheet per DUT config group -- same Mean/Std/
   Min/Max/USL/LSL (+ Delta) shape as the existing pooled "BT_TX Summary"
   sheet (kept as-is, untouched), but Mean/Std/Min/Max computed from ONLY
   that group's DUTs; USL/LSL are the SAME as the pooled summary row for
   that Test Item (spec doesn't depend on grouping) -- pulled straight off
   the already-built "BT_TX Summary" sheet's own cells rather than
   re-deriving a spec lookup a second time.
2. REPLACE the existing flat "Raw Values" sheet with a 3-level-header
   version grouped Bench/ATE/Delta -> config group -> individual DUT serial
   -- same shape as Harald's own grouped Raw Values.
3. Sheet order: "BT_TX Summary", then one "<Config> Summary" sheet per
   sorted config group, then "Raw Values" last -- NOT hardcoded to any
   fixed group count. Harald's own docstring (update_harald_260820_summary_
   by_vendor.py, module doc point 1) documents a real past bug from
   hardcoding an exact group count (was "exactly 2 vendors", broke on a
   100-DUT/4-group pull); this script derives the group set purely from
   whatever Config labels 40pcs_config.txt actually contains, sorted for
   determinism, and never assumes 2.

Why config, not vendor-folder: Harald derives group membership from Bench's
own folder structure (subfolder name = group). BT_TX's Bench folder is
nested by DUT instead (Bench/#NN_<serial>/...), one folder per DUT -- there
is no folder-based way to recover group membership. Per user confirmation
("TX에서는 벤더를 구분할만한 데가 없으니 40pcs_config.txt 이걸로 확인하라는
말이었어"), group membership instead comes from an external mapping file
the user maintains by hand: 40pcs_config.txt (same directory convention as
every other shared BT_TX file -- see CONFIG_PATH below), columns Num/2DID/
Config, where 2DID is the same DUT-serial string the Bench-filename-embedded
serial and the ATE CSV header already use, and Config is the group label
(currently SNKJ/DNNJ, 25/15 of the 40 DUTs -- verified from the real file,
not assumed 20/20).

Reuses add_bt_tx_raw_values_sheet.py's 2026-08-25 refactor (collect_raw_rows/
parse_spec_tables/flatten_rows/write_flat_raw_values_sheet) for the actual
per-DUT Bench/ATE data and the flat-sheet fallback -- never re-derives a
stat or duplicates that traversal.
"""
from __future__ import annotations

import datetime
import statistics
import sys
from pathlib import Path as _Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

_THIS_DIR = _Path(__file__).resolve().parent
_BASE_DIR = _THIS_DIR.parent
sys.path.insert(0, str(_BASE_DIR))
sys.path.insert(0, str(_THIS_DIR))

import paths as bt_tx_paths
import xlsx_summary_writer as xw
import add_bt_tx_raw_values_sheet as raw_sheet

PATHS = bt_tx_paths.get_paths("BT_TX")
REPORT_DATE = datetime.date.today().isoformat()
SUMMARY_PATH = PATHS.report_dir / f"{REPORT_DATE}_BT_TX_Summary.xlsx"

# Same one-level-up convention every other BT_TX script uses to locate a file
# shared with the rest of the workspace (paths.py, config.xml, ...) -- this
# resolves to the root copy when run from root BT_TX/, and to UI's copy when
# run from UI/BT_TX/, since _BASE_DIR is this script's own parent directory's
# parent (root, or UI/) in either location.
CONFIG_PATH = _BASE_DIR / "40pcs_config.txt"


def parse_config_map(path: _Path) -> tuple[dict, dict]:
    """40pcs_config.txt (tab-or-whitespace-separated, header row Num/2DID/
    Config) -> ({serial: config_label}, {config_label: [serial, ...]}
    sorted). The Num column ("#NN") is ignored -- it's redundant with the
    same index already recovered from Bench's own filenames elsewhere in
    this pipeline. Group set is whatever Config labels actually appear in
    the file, never hardcoded (see module docstring)."""
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    serial_to_config: dict = {}
    if not lines:
        return serial_to_config, {}

    header = lines[0].split()
    header_lower = [h.lower() for h in header]
    try:
        idx_2did = header_lower.index("2did")
        idx_config = header_lower.index("config")
    except ValueError:
        return serial_to_config, {}

    for line in lines[1:]:
        parts = line.split()
        if len(parts) <= max(idx_2did, idx_config):
            continue
        serial = parts[idx_2did]
        config = parts[idx_config]
        serial_to_config[serial] = config

    config_to_serials: dict = {}
    for serial, config in serial_to_config.items():
        config_to_serials.setdefault(config, []).append(serial)
    config_to_serials = {c: sorted(s) for c, s in config_to_serials.items()}
    return serial_to_config, config_to_serials


def spec_by_name_from_summary_sheet(wb) -> dict:
    """{Test Item name: (bench_usl, bench_lsl, ate_usl, ate_lsl)}, read
    straight off the already-built "BT_TX Summary" sheet's own cells --
    columns 6/7 (Bench USL/LSL) and 12/13 (ATE USL/LSL), per xlsx_summary_
    writer.BENCH_SUB/ATE_SUB order (Mean,Std,Min,Max,USL,LSL) and that
    sheet's own 2-row merged header (data starts row 3). Reused rather than
    re-deriving a spec lookup a second time -- spec doesn't depend on
    grouping, so the pooled summary's own USL/LSL is exactly what every
    "<Config> Summary" row should show too."""
    ws = wb["BT_TX Summary"]
    out = {}
    for r in range(3, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if not name:
            continue
        bench_usl = ws.cell(row=r, column=6).value
        bench_lsl = ws.cell(row=r, column=7).value
        ate_usl = ws.cell(row=r, column=12).value
        ate_lsl = ws.cell(row=r, column=13).value
        out[name] = (bench_usl, bench_lsl, ate_usl, ate_lsl)
    return out


def _stats(values):
    if not values:
        return None
    return (statistics.mean(values), min(values), max(values),
            statistics.stdev(values) if len(values) > 1 else 0.0)


def _r(v):
    return round(v, 4) if isinstance(v, (int, float)) else v


def _stats_cells(stats):
    if not stats:
        return [None, None, None, None]
    mean, vmin, vmax, std = stats
    return [_r(mean), _r(std), _r(vmin), _r(vmax)]


def build_config_summary_rows(dict_rows, config_serials: set, spec_by_name: dict) -> list:
    """Same shape as build_bt_tx_summary_csv's own rows (name, Bench Mean/
    Std/Min/Max/USL/LSL, ATE Mean/Std/Min/Max/USL/LSL) + 4 Delta columns,
    but Mean/Std/Min/Max recomputed from ONLY `config_serials`'s values out
    of each row's (name, bench_by_serial, ate_by_serial) -- USL/LSL kept as
    the pooled summary's own (spec_by_name), never re-derived."""
    rows = []
    for name, bench_by_serial, ate_by_serial in dict_rows:
        bench_vals = [v for s, v in bench_by_serial.items() if s in config_serials]
        ate_vals = [v for s, v in ate_by_serial.items() if s in config_serials]
        bench_stats = _stats(bench_vals)
        ate_stats = _stats(ate_vals) if ate_vals else None
        bench_usl, bench_lsl, ate_usl, ate_lsl = spec_by_name.get(name, (None, None, None, None))
        bench_cells = _stats_cells(bench_stats) + [bench_usl, bench_lsl]
        ate_cells = _stats_cells(ate_stats) + [ate_usl, ate_lsl]
        deltas = [xw.delta(bench_cells[i], ate_cells[i]) for i in range(4)]
        rows.append([name, *bench_cells, *ate_cells, *deltas])
    return rows


def write_config_summary_sheet(wb, sheet_name: str, rows: list, n_cols: int = 17) -> None:
    """Creates/replaces one "<Config> Summary" sheet -- same merged 2-row
    header (Bench/ATE/Delta -> Mean/Std/Min/Max[/USL/LSL]) as xlsx_summary_
    writer.write_summary_xlsx, mirrors update_harald_260820_summary_by_
    vendor.py's write_vendor_summary_sheet exactly. Does NOT save `wb`."""
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    ws.cell(row=1, column=1, value="Test Item")
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    col = 2
    for group_label, subs in (("Bench", xw.BENCH_SUB), ("ATE", xw.ATE_SUB), ("Delta", xw.DELTA_SUB)):
        start = col
        for sub in subs:
            ws.cell(row=2, column=col, value=sub)
            col += 1
        ws.cell(row=1, column=start, value=group_label)
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=col - 1)

    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=2, max_col=n_cols):
        for cell in row:
            cell.font = header_font
            cell.alignment = center
    for i in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14

    for r, row in enumerate(rows, start=3):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    ws.freeze_panes = "A3"


def write_grouped_raw_values_sheet(wb, dict_rows, config_order, config_to_serials) -> None:
    """Creates/replaces "Raw Values" with the 3-level header (Bench/ATE/
    Delta -> config group -> serial) -- does NOT save `wb`. Mirrors
    update_harald_260820_summary_by_vendor.py's own Raw Values rebuild
    exactly, substituting BT_TX's (name, bench_by_serial, ate_by_serial)
    dict_rows for Harald's job-based lookup."""
    rv_rows = []
    for name, bench_by_serial, ate_by_serial in dict_rows:
        row = [name]
        for group in ("bench", "ate", "delta"):
            for config in config_order:
                for serial in config_to_serials[config]:
                    b = bench_by_serial.get(serial)
                    a = ate_by_serial.get(serial)
                    if group == "bench":
                        row.append(b)
                    elif group == "ate":
                        row.append(a)
                    else:
                        row.append(xw.delta(b, a))
        rv_rows.append(row)

    sheet_name = "Raw Values"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.cell(row=1, column=1, value="Test Item")
    ws.merge_cells(start_row=1, start_column=1, end_row=3, end_column=1)

    col = 2
    for group_label in ("Bench", "ATE", "Delta"):
        group_start = col
        for config in config_order:
            config_start = col
            for serial in config_to_serials[config]:
                ws.cell(row=3, column=col, value=serial)
                col += 1
            ws.cell(row=2, column=config_start, value=config)
            ws.merge_cells(start_row=2, start_column=config_start, end_row=2, end_column=col - 1)
        ws.cell(row=1, column=group_start, value=group_label)
        ws.merge_cells(start_row=1, start_column=group_start, end_row=1, end_column=col - 1)
    n_cols = col - 1

    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=3, max_col=n_cols):
        for cell in row:
            cell.font = header_font
            cell.alignment = center
    for i in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 12

    for r, row in enumerate(rv_rows, start=4):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    ws.freeze_panes = "B4"


def fallback_flat_raw_values(wb, num_to_serial) -> None:
    """Builds the OLD flat (ungrouped, 2-level Bench/ATE -> serial header)
    Raw Values sheet -- the graceful-degradation path when 40pcs_config.txt
    is missing or incomplete for this Bench pull. No "<Config> Summary"
    sheets are written in this path either (no group info to build them
    from) -- the workbook ends up with just "BT_TX Summary" + this flat
    "Raw Values", same as before this feature existed. Does NOT save `wb`."""
    ate_serials = raw_sheet.collect_ate_serials()
    all_serials = sorted(set(num_to_serial.values()) | ate_serials)
    devm_spec_table, freqacc_spec_table = raw_sheet.parse_spec_tables()
    dict_rows = raw_sheet.collect_raw_rows(num_to_serial, devm_spec_table, freqacc_spec_table)
    flat_rows = raw_sheet.flatten_rows(dict_rows, all_serials)
    raw_sheet.write_flat_raw_values_sheet(wb, flat_rows, all_serials)
    print(f"  wrote flat Raw Values: {len(flat_rows)} row(s) x {1 + 2 * len(all_serials)} col(s)")


def main() -> None:
    print("Building #NN -> DUT serial map from the Bench folder...")
    num_to_serial = raw_sheet.build_num_to_serial(PATHS.bench_dir)
    if not num_to_serial:
        print("Bench folder has no embedded DUT serials (old flat layout) -- "
              "nothing to group by config; skipping Raw Values sheet "
              "entirely (same as add_bt_tx_raw_values_sheet.py's own rule).")
        return
    print(f"  {len(num_to_serial)} DUT(s) mapped")

    wb = openpyxl.load_workbook(SUMMARY_PATH)

    if not CONFIG_PATH.exists():
        print(f"{CONFIG_PATH} not found -- falling back to the flat "
              f"(ungrouped) Raw Values sheet, no <Config> Summary sheets.")
        fallback_flat_raw_values(wb, num_to_serial)
        wb.save(SUMMARY_PATH)
        print(f"Saved: {SUMMARY_PATH}. Sheets: {wb.sheetnames}")
        return

    print(f"Parsing DUT config map ({CONFIG_PATH.name})...")
    serial_to_config, config_to_serials = parse_config_map(CONFIG_PATH)
    if not serial_to_config:
        print(f"{CONFIG_PATH.name} has no usable 2DID/Config columns -- "
              f"falling back to the flat (ungrouped) Raw Values sheet, no "
              f"<Config> Summary sheets.")
        fallback_flat_raw_values(wb, num_to_serial)
        wb.save(SUMMARY_PATH)
        print(f"Saved: {SUMMARY_PATH}. Sheets: {wb.sheetnames}")
        return

    bench_serials = set(num_to_serial.values())
    missing = sorted(s for s in bench_serials if s not in serial_to_config)
    if missing:
        print(f"{len(missing)} Bench DUT serial(s) have no entry in "
              f"{CONFIG_PATH.name} (e.g. {missing[:5]}) -- falling back to "
              f"the flat (ungrouped) Raw Values sheet rather than guessing "
              f"a group (no <Config> Summary sheets either).")
        fallback_flat_raw_values(wb, num_to_serial)
        wb.save(SUMMARY_PATH)
        print(f"Saved: {SUMMARY_PATH}. Sheets: {wb.sheetnames}")
        return

    config_order = sorted(config_to_serials)
    for c in config_order:
        print(f"  {c}: {len(config_to_serials[c])} DUT(s)")

    print("Parsing spec workbooks (needed for DEVM/FreqAcc's own outlier/"
          "range windows, to keep this sheet's row set identical to the "
          "summary sheet's)...")
    devm_spec_table, freqacc_spec_table = raw_sheet.parse_spec_tables()

    dict_rows = raw_sheet.collect_raw_rows(num_to_serial, devm_spec_table, freqacc_spec_table)
    print(f"Total rows: {len(dict_rows)}")

    print("Reading USL/LSL off the existing 'BT_TX Summary' sheet (spec "
          "doesn't depend on grouping)...")
    spec_by_name = spec_by_name_from_summary_sheet(wb)

    for config in config_order:
        print(f"Building '{config} Summary' sheet...")
        rows = build_config_summary_rows(dict_rows, set(config_to_serials[config]), spec_by_name)
        write_config_summary_sheet(wb, f"{config} Summary", rows)

    print("Rebuilding 'Raw Values' sheet (3-level header: Bench/ATE/Delta -> config -> serial)...")
    write_grouped_raw_values_sheet(wb, dict_rows, config_order, config_to_serials)

    # Sheet order: pooled summary (untouched), one "<Config> Summary" per
    # sorted group, then grouped Raw Values last -- never hardcoded to a
    # fixed group count (see module docstring).
    desired_order = ["BT_TX Summary", *(f"{c} Summary" for c in config_order), "Raw Values"]
    wb._sheets.sort(key=lambda s: desired_order.index(s.title) if s.title in desired_order else len(desired_order))

    print(f"\nSaving {SUMMARY_PATH}...")
    wb.save(SUMMARY_PATH)
    print(f"Saved. Sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
