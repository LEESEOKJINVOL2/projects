"""BT TX EVM (preamble, dB) distribution charts.

For every unique combination of the 5 config columns
  cfg-carrier_01-packet_type, -pa_supply, -pa_slices, -pa_gain, -dig_gain
plot avg_evm_preamble_aver-dB (Y) against cfg-carrier_01-frequency-MHz (X),
overlaying all 40 DUTs as separate line series (the distribution).

Two outputs per combo:
  * Excel workbook: chart on top, extracted data table below.
  * PNG image.

File name encodes the combo, e.g.
  packet_type(1DH5)_pa_supply(0)_pa_slices(1)_pa_gain(0)_dig_gain(64)
(the "cfg-carrier_01-" prefix is stripped). Chart title is prefixed with
"BT_TX_".
"""

from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import cm

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Font, PatternFill

import sys
from pathlib import Path as _Path
sys.path.insert(0, str(_Path(__file__).resolve().parent.parent.parent))
import paths as bt_tx_paths

PATHS = bt_tx_paths.get_paths("BT_TX")
from openpyxl.utils import get_column_letter


KEY_COLS = [
    "cfg-carrier_01-packet_type",
    "cfg-carrier_01-pa_supply",
    "cfg-carrier_01-pa_slices",
    "cfg-carrier_01-pa_gain",
    "cfg-carrier_01-dig_gain",
]
X_COL = "cfg-carrier_01-frequency-MHz"
Y_COL = "avg_evm_preamble_aver-dB"

DUT_RE = re.compile(r"_#(\d+)_")


def strip_prefix(col: str) -> str:
    return col.replace("cfg-carrier_01-", "")


def combo_name(key: tuple[str, ...]) -> str:
    parts = [f"{strip_prefix(col)}({val})" for col, val in zip(KEY_COLS, key)]
    return "_".join(parts)


def safe_filename(name: str) -> str:
    return re.sub(r'[\\/*?:\[\]<>|"]', "_", name)


def dut_from_filename(path: Path) -> int | None:
    m = DUT_RE.search(path.name)
    return int(m.group(1)) if m else None


def to_float(text: str):
    text = (text or "").strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def read_all(bt_tx_dir: Path):
    """Return (duts, data) where data[key][dut][freq] = avg_evm_preamble_aver-dB."""
    data: dict[tuple, dict[int, dict[float, float]]] = defaultdict(
        lambda: defaultdict(dict)
    )
    duts: set[int] = set()
    files = sorted(bt_tx_dir.glob("*.csv"))
    for path in files:
        dut = dut_from_filename(path)
        if dut is None:
            print(f"  WARN: no DUT number in {path.name}, skipped")
            continue
        duts.add(dut)
        with path.open(encoding="utf-8-sig", newline="") as fh:
            reader = csv.DictReader(fh)
            for row in reader:
                freq = to_float(row.get(X_COL, ""))
                if freq is None:
                    continue
                evm = to_float(row.get(Y_COL, ""))
                if evm is None:
                    continue
                key = tuple((row.get(c, "") or "").strip() for c in KEY_COLS)
                data[key][dut][freq] = evm
        print(f"  read {path.name} (DUT #{dut})")
    return sorted(duts), data


def freq_axis(data_for_combo: dict[int, dict[float, float]]) -> list[float]:
    freqs: set[float] = set()
    for dut_map in data_for_combo.values():
        freqs.update(dut_map.keys())
    return sorted(freqs)


def color_hex(i: int, n: int) -> str:
    r, g, b, _ = cm.get_cmap("hsv")(i / max(1, n))
    return f"{int(r*255):02X}{int(g*255):02X}{int(b*255):02X}"


def write_png(png_path: Path, title: str, freqs, duts, combo_data) -> None:
    fig, ax = plt.subplots(figsize=(14, 7))
    for i, dut in enumerate(duts):
        ys = [combo_data.get(dut, {}).get(f) for f in freqs]
        ax.plot(freqs, ys, linewidth=1.0, color=f"#{color_hex(i, len(duts))}",
                label=f"#{dut}")
    ax.set_title(title, fontsize=9)
    ax.set_xlabel("Frequency (MHz)")
    ax.set_ylabel("EVM (dB)")
    ax.grid(True, which="major", color="0.8", linewidth=0.5)
    ax.set_xticks(freqs)
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.1), ncol=10,
              fontsize=6, frameon=False)
    fig.tight_layout()
    fig.savefig(png_path, dpi=130, bbox_inches="tight")
    plt.close(fig)


def write_xlsx(xlsx_path: Path, title: str, freqs, duts, combo_data) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "EVM_dB"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    chart_row_span = 30
    data_header_row = 1 + chart_row_span

    # Data table: frequency column + one column per DUT.
    ws.cell(data_header_row, 1, "Frequency (MHz)")
    for j, dut in enumerate(duts, 2):
        ws.cell(data_header_row, j, f"#{dut}")
    for r, freq in enumerate(freqs, data_header_row + 1):
        ws.cell(r, 1, freq)
        for j, dut in enumerate(duts, 2):
            v = combo_data.get(dut, {}).get(freq)
            if v is not None:
                ws.cell(r, j, v)
    for cell in ws[data_header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions["A"].width = 16
    for j in range(2, len(duts) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 9

    last_row = data_header_row + len(freqs)
    chart = LineChart()
    chart.title = title
    chart.style = 10
    chart.grouping = "standard"
    chart.height = 15
    chart.width = 34
    chart.y_axis.title = "EVM (dB)"
    chart.x_axis.title = "Frequency (MHz)"
    chart.y_axis.delete = False
    chart.x_axis.delete = False
    chart.y_axis.numFmt = "0.00"
    chart.x_axis.tickLblPos = "low"
    chart.y_axis.majorGridlines = ChartLines()
    chart.y_axis.majorGridlines.spPr = GraphicalProperties()
    chart.y_axis.majorGridlines.graphicalProperties.line.solidFill = "808080"
    chart.y_axis.majorGridlines.graphicalProperties.line.width = 12700
    chart.legend.position = "b"
    chart.legend.overlay = False

    cats = Reference(ws, min_col=1, min_row=data_header_row + 1, max_row=last_row)
    for j, dut in enumerate(duts):
        col = j + 2
        ref = Reference(ws, min_col=col, min_row=data_header_row, max_row=last_row)
        chart.add_data(ref, titles_from_data=True)
        series = chart.series[-1]
        series.graphicalProperties.line.solidFill = color_hex(j, len(duts))
        series.graphicalProperties.line.width = 10000
        series.marker.symbol = "none"
    chart.set_categories(cats)
    ws.add_chart(chart, "A1")

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--bt-tx-dir", type=Path, default=PATHS.bench_dir)
    p.add_argument("--out-dir", type=Path, default=PATHS.result_xlsx_dir / "evm_dB")
    p.add_argument("--png-dir", type=Path, default=PATHS.result_png_dir / "evm_dB")
    p.add_argument("--limit", type=int, default=0,
                   help="Generate only the first N combos (0 = all).")
    return p.parse_args()


def main() -> None:
    args = parse_args()
    print("Reading BT_Tx CSVs...")
    duts, data = read_all(args.bt_tx_dir)
    print(f"DUTs: {len(duts)}  combos: {len(data)}")

    keys = list(data.keys())
    if args.limit > 0:
        keys = keys[: args.limit]
        print(f"Limiting to first {len(keys)} combo(s)")

    args.out_dir.mkdir(parents=True, exist_ok=True)
    args.png_dir.mkdir(parents=True, exist_ok=True)

    for n, key in enumerate(keys, 1):
        name = combo_name(key)
        fname = safe_filename(name)
        title = f"BT_TX_{name}"
        freqs = freq_axis(data[key])
        write_xlsx(args.out_dir / f"{fname}.xlsx", title, freqs, duts, data[key])
        write_png(args.png_dir / f"{fname}.png", title, freqs, duts, data[key])
        if n % 50 == 0 or n == len(keys):
            print(f"  [{n}/{len(keys)}] {fname}")

    print(f"\nDone. XLSX -> {args.out_dir}")
    print(f"      PNG  -> {args.png_dir}")


if __name__ == "__main__":
    main()
