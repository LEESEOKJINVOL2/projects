"""Parse the 'TX Power Specification' sheet
(Harald_P2_Test_Coverage_v4.4_JSCK.xlsx) into a flat lookup table:
  spec[(band_bucket, packet_type, pa_supply)] = (LSL, USL)

2026-08-14: replaces the old 'BT Performance Specification' sheet parse
(NPI_TestCoverage_v0.4.xlsx) for Power only -- DEVM/ACP/FreqAcc still read
the old workbook via their own *_spec_lookup.py modules, see paths.py's
power_spec_workbook docstring for why this is a per-project override
rather than swapping the shared spec_workbook.

Sheet layout (rows 4-42; data starts at column B, column A is blank on
every row -- same "column A unused" convention as this workbook's other
new sheets):
  col B = Band (only set on the first row of a band block -- forward-
    filled: "B0" then "B1,B2, B4, B5, B6, B7"). Unlike the old workbook,
    THIS spec does not distinguish B1 (5150-5250MHz) from B2/B4-B7 -- every
    5/6GHz frequency shares one row block, confirmed by the single
    "B1,B2, B4, B5, B6, B7" label spanning all of them (2026-08-14).
  col C = Packet Type -- a "/"-joined list of packet-type tokens that all
    share this row's limit (e.g. "HDT2/HDT3/HDRPS2/HDRPM4/HDRPM6/HDRPM8").
    Exploded directly into individual lookup keys below -- no intermediate
    "parameter group" indirection needed like the old sheet's PACKET_TO_
    GROUP dict, since this sheet already lists every packet type per row.
    A handful of tokens use legacy protocol-name labels instead of the
    bench CSV's own packet_type spelling (BDR/EDR2/EDR3/HDR4/HDR8 vs
    1DH5/2DH5/3DH5/4DH5/8DH5) -- translated via LEGACY_LABEL_TO_PACKET_TYPE,
    everything else in the column already matches the bench CSV's
    packet_type strings literally (LE1M, HDT4, HDRPS2, UHDR48, etc.).
  col D = VDD_HPA [V] -- 1.5/1.2/0.73, maps 1:1 to the bench CSV's
    cfg-carrier_01-pa_supply index 0/1/2 per this same workbook's "Tx Power
    Table Condition Specification" footer table (rows 45-49): HPA=1.5 ->
    supply 0, 1.2 -> 1, 0.73 -> 2.
  col E = LSL, col F = Typ., col G = USL, col H = Units.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

SHEET_NAME = "TX Power Specification"
ROW_START = 4
ROW_END = 42
COL_BAND = 2
COL_PACKET_TYPES = 3
COL_PA_SUPPLY_V = 4
COL_LSL = 5
COL_USL = 7

PA_SUPPLY_VOLT_TO_INDEX = {"1.5": "0", "1.2": "1", "0.73": "2"}

LEGACY_LABEL_TO_PACKET_TYPE = {
    "BDR": "1DH5",
    "EDR2": "2DH5",
    "EDR3": "3DH5",
    "HDR4": "4DH5",
    "HDR8": "8DH5",
}


def band_bucket_of(freq_mhz: float) -> str:
    """B0 = 2.4GHz, everything else (5/6GHz) shares one bucket -- see
    module docstring, this sheet doesn't split out B1 the way the old one
    did."""
    return "B0" if freq_mhz < 3000 else "B1_B2_B4_B5_B6_B7"


def parse_spec(xlsx_path: Path):
    """Return spec[(band_bucket, packet_type, pa_supply)] = (LSL, USL)."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[SHEET_NAME]

    spec: dict[tuple, tuple[float, float]] = {}
    current_band_bucket = None

    for r in range(ROW_START, ROW_END + 1):
        band_cell = ws.cell(r, COL_BAND).value
        if band_cell:
            band_str = str(band_cell).strip()
            current_band_bucket = "B0" if band_str == "B0" else "B1_B2_B4_B5_B6_B7"

        pkt_cell = ws.cell(r, COL_PACKET_TYPES).value
        volt_cell = ws.cell(r, COL_PA_SUPPLY_V).value
        lsl = ws.cell(r, COL_LSL).value
        usl = ws.cell(r, COL_USL).value
        if not pkt_cell or volt_cell is None or lsl is None or usl is None or current_band_bucket is None:
            continue

        pa_supply = PA_SUPPLY_VOLT_TO_INDEX.get(f"{float(volt_cell):g}")
        if pa_supply is None:
            continue

        tokens = {t.strip() for t in str(pkt_cell).split("/") if t.strip()}
        for tok in tokens:
            packet_type = LEGACY_LABEL_TO_PACKET_TYPE.get(tok, tok)
            spec[(current_band_bucket, packet_type, pa_supply)] = (float(lsl), float(usl))

    return spec


def lookup(spec: dict, packet_type: str, freq_mhz: float, pa_supply: str):
    """Return (LSL, USL) or None if this packet/frequency/supply has no spec."""
    band_bucket = band_bucket_of(freq_mhz)
    return spec.get((band_bucket, packet_type, pa_supply))


if __name__ == "__main__":
    spec = parse_spec(Path(__file__).resolve().parent.parent / "Harald_P2_Test_Coverage_v4.4_JSCK.xlsx")
    print(f"total spec entries: {len(spec)}")
    tests = [
        (("B0", "1DH5", "0"), (14.5, 16.5)),
        (("B0", "HDT4", "2"), (4.4, 6.4)),
        (("B1_B2_B4_B5_B6_B7", "LE1M", "1"), (10.7, 13.2)),
        (("B1_B2_B4_B5_B6_B7", "UHDR48", "0"), (8, 10.5)),
    ]
    for key, expected in tests:
        got = spec.get(key)
        print(key, "->", got, "  expected:", expected, "OK" if got == expected else "MISMATCH")

    print()
    print("lookup(1DH5, 2441, '0') ->", lookup(spec, "1DH5", 2441, "0"))
    print("lookup(3DH5, 5787, '1') ->", lookup(spec, "3DH5", 5787, "1"))
    print("lookup(HDRPS2, 6234, '0') ->", lookup(spec, "HDRPS2", 6234, "0"))
