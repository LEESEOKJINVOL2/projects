"""Shared XLSX-writing helper for per-project Bench/ATE/Delta summary
scripts (BT_TX/build_bt_tx_xlsx_summary.py, BT_RX/build_bt_rx_xlsx_summary.py,
Cal/build_cal_xlsx_summary.py) -- added 2026-08-15 per user request for a
Mean/Std/Min/Max (+ delta) summary in .xlsx form. Lives at the workspace
root alongside paths.py/png_jobs.py since it's generic formatting, not
project-specific data logic -- each caller still owns its own row-building
(reusing that project's own Bench/ATE reader modules), this only writes the
resulting rows to a merged-header worksheet.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

BENCH_SUB = ["Mean", "Std", "Min", "Max", "USL", "LSL"]
ATE_SUB = ["Mean", "Std", "Min", "Max", "USL", "LSL"]
DELTA_SUB = ["Mean", "Std", "Min", "Max"]


def delta(a, b):
    """|a-b|, or None if either side is missing. Per user instruction
    2026-08-15 ("전부 abs 처리해주면 되고" -- just abs() everything): every
    delta column is an absolute difference, never signed. Only meant for
    Mean/Std/Min/Max -- USL/LSL never gets a delta column since it's the
    same physical spec limit on both sides by construction, not an
    independently measured value."""
    if a is None or b is None:
        return None
    return abs(a - b)


def write_summary_xlsx(rows: list, out_path: Path, sheet_title: str) -> None:
    """rows: [[Test Item, Bench Mean/Std/Min/Max/USL/LSL, ATE Mean/Std/Min/
    Max/USL/LSL, Delta Mean/Std/Min/Max], ...] -- 17 columns total. Writes a
    merged 2-row header (Bench/ATE/Delta group label spanning its own
    sub-columns) -- the CSV-format sibling scripts (build_*_summary_csv.py)
    describe this as the original xlsx design, before a 2026-07-22 customer
    request replaced it with a flat-header CSV; this revives it for the
    xlsx format the user now wants, unrelated to that CSV decision."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    ws.cell(row=1, column=1, value="Test Item")
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)

    col = 2
    for group_label, subs in (("Bench", BENCH_SUB), ("ATE", ATE_SUB), ("Delta", DELTA_SUB)):
        start = col
        for sub in subs:
            ws.cell(row=2, column=col, value=sub)
            col += 1
        ws.cell(row=1, column=start, value=group_label)
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=col - 1)
    n_cols = col - 1

    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=2, max_col=n_cols):
        for cell in row:
            cell.font = header_font
            cell.alignment = center
    for i in range(1, n_cols + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 14

    for r, row in enumerate(rows, start=3):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)

    ws.freeze_panes = "A3"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
