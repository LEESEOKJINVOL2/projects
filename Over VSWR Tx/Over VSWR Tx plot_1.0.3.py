import os
import pandas as pd
import matplotlib.pyplot as plt
from tkinter import Tk, Button, filedialog, messagebox
from pptx import Presentation
from pptx.util import Cm

def plot(file_path):
    png_path = os.getcwd() + "\\PNG"  # delete all files in directory at first
    if os.path.exists(png_path):
        for file in os.scandir(png_path):
            os.remove(file)
        print("delete all files in PNG directory")
    else:
        print("delete error")

    try:
        df = pd.read_csv(file_path)
    except FileNotFoundError:
        messagebox.showerror("Error", f"'{file_path}' doesn`t exist. Please check directory")
        return
    
    if df.empty:
        messagebox.showwarning("Warning", "CSV file is Empty")
        return
    
    bands = df['BAND'].unique()  # Load valid data
    pkt_types = df['PKT_TYPE'].unique()
    hpa_types = df['HPA'].unique()  # Add HPA unique values
    output_dir = 'png'
    os.makedirs(output_dir, exist_ok=True)

    for pkt in pkt_types:  # PKT_TYPE(DH5, 8-DH5, 8-HDRP)
        pkt_data = df[df['PKT_TYPE'] == pkt]
        
        for hpa in hpa_types:  # Filter by HPA type
            hpa_data = pkt_data[pkt_data['HPA'] == hpa]
            
            for band in bands:  # BAND
                band_data = hpa_data[hpa_data['BAND'] == band]
                channels = band_data['CH'].unique()
                
                for ch in channels:  # CH
                    filtered_data = band_data[band_data['CH'] == ch]
                    
                    plt.figure(figsize=(12, 24)) # png size
                    for vswr in [1, 2, 3]:  # VSWR
                        vswr_data = filtered_data[filtered_data['VSWR'] == vswr]
                        
                        if not vswr_data.empty:
                            plt.plot(vswr_data['PHASE'], vswr_data['VALUE'], marker='o', label=f'VSWR {vswr}', linewidth = 10, markersize = 20) # line style
                    
                    base_filename = os.path.splitext(os.path.basename(file_path))[0]  # csv file name

                    plt.title(f'Band = {band}\n', fontsize = 60)  # plot title
                    plt.xlabel('phase [degree]', fontsize = 40)  # x label
                    if band == 0:
                        plt.ylabel('Fc [dBm]', fontsize = 40)  # y label
                    else:
                        # plt.gca().yaxis.set_visible([]) # y axis delete
                        ax = plt.gca()
                        ax.tick_params(axis='y', labelleft=False)
                    plt.xticks(range(0, 360, 30), rotation=45, fontsize = 40) # range 0 to 360, span 30

                    if pkt == 'DH5':  # y span per each packet & HPA
                        if hpa == 730:
                            plt.yticks(range(4, 14, 1), fontsize = 40)
                        elif hpa == 1200:
                            plt.yticks(range(9, 18, 1), fontsize = 40)
                        elif hpa == 1500:
                            plt.yticks(range(11, 20, 1), fontsize = 40)
                    else:
                        if hpa == 730:
                            plt.yticks(range(1, 10, 1), fontsize = 40)
                        elif hpa == 1200:
                            plt.yticks(range(6, 15, 1), fontsize = 40)
                        elif hpa == 1500:
                            plt.yticks(range(8, 17, 1), fontsize = 40)

                    if band == 7:
                        plt.legend(ncol=3, loc = 'upper right', fontsize = 25) # legend option
                    plt.grid(True)
                    
                    filename = os.path.join(output_dir, f'{base_filename}_{pkt}_HPA {hpa}_{ch}_Band {band}.png') # png file name
                    plt.savefig(filename, bbox_inches='tight')
                    plt.close()

    messagebox.showinfo("Success", "Successfully saved png files")

def ppt():
    output_dir = 'png'
    png_files = [f for f in os.listdir(output_dir) if f.endswith('.png')]
    if not png_files:
        messagebox.showwarning("Warning", "No png file, Please make png files")
        return
    
    presentation = Presentation('JCET_Format.pptx') # JCET ppt templet
    groups = {}
    
    for png_file in png_files:
        parts = png_file.split('_') # ppt title
        title = "_".join(parts[:4])
        
        if title not in groups:
            groups[title] = []
        
        groups[title].append(png_file)

    for title, files in groups.items():
        slide = presentation.slides.add_slide(presentation.slide_layouts[1]) # using 2nd ppt slide
        
        for idx, png_file in enumerate(files):
            img_path = os.path.join(output_dir, png_file)

            width = Cm(4.75) # png width 
            height = Cm(10) # png height

            left = Cm(0.32 + (idx * 4.75))
            top = Cm(4.52)

            slide.shapes.add_picture(img_path, left, top, width=width, height=height)

        title_box = slide.shapes.title
        title_box.text = title

    save_path = filedialog.asksaveasfilename(defaultextension=".pptx", filetypes=[("PowerPoint files", "*.pptx")], title="Save PPT File", initialfile=f"{parts[0]}.pptx")
    
    if save_path:
        presentation.save(save_path)
        messagebox.showinfo("Success", f"Successfully saved as '{save_path}'")

def open_file():
    file_path = filedialog.askopenfilename(title="Select CSV File", filetypes=[("CSV Files", "*.csv")])
    if file_path:
        plot(file_path)

root = Tk()
root.title("Over VSWR Plot_1.0.3")
root.geometry("300x200")

open_button = Button(root, text="Open CSV", command=open_file)
open_button.pack(pady=20)

ppt_button = Button(root, text="Make PPT", command=ppt)
ppt_button.pack(pady=20)

quit_button = Button(root, text="Quit", command=root.quit)
quit_button.pack(pady=20)

root.mainloop()